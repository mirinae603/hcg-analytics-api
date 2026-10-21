"""Streaming Excel readers (.xlsx and .xlsb) and cell-value normalisers.

The monthly sales sheets are ~880k rows x 33 columns, so they are never loaded whole:
.xlsx goes through openpyxl read_only/data_only, .xlsb through pyxlsb, one row at a time
(the pattern proven on these files in hcg-work/convert.py: 120-180 s per sheet).
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Iterator, Optional

from app.etl.formulary.classify import is_blank, is_error_value, norm_code, norm_text

EXCEL_BASE = datetime(1899, 12, 30)
HEADER_SCAN_ROWS = 10


class Workbook:
    """Read-only view of one .xlsx/.xlsb file: sheet names and a row iterator."""

    def __init__(self, path):
        self.path = Path(path)
        self.kind = self.path.suffix.lower()
        self._sheets: dict = {}
        if self.kind == ".xlsb":
            from pyxlsb import open_workbook   # a hard requirement: no silent empty reads
            self._wb = open_workbook(str(self.path))
            self.sheet_names = list(self._wb.sheets)
        elif self.kind in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            self._wb = load_workbook(str(self.path), read_only=True, data_only=True)
            self.sheet_names = list(self._wb.sheetnames)
        else:
            raise ValueError(f"not an Excel workbook: {self.path.name}")

    def rows(self, sheet: str) -> Iterator[list]:
        if self.kind == ".xlsb":
            # pyxlsb unpacks a sheet into a temp file on every get_sheet(); keep it open so
            # the header scan and the full read unpack the 880k-row sheet only once
            if sheet not in self._sheets:
                self._sheets[sheet] = self._wb.get_sheet(sheet)
            for r in self._sheets[sheet].rows():
                yield [c.v for c in r]
        else:
            for r in self._wb[sheet].iter_rows(values_only=True):
                yield r

    def close(self):
        for s in self._sheets.values():
            try:
                s.close()
            except Exception:
                pass
        self._sheets.clear()
        try:
            self._wb.close()
        except Exception:
            pass

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self.close()


def header_key(v) -> str:
    return norm_text(v)


def find_header(rows: Iterator[list], must_have, scan: int = HEADER_SCAN_ROWS):
    """(row index, normalised header list) of the first row within `scan` rows that
    contains every name in must_have; (None, None) if there is none."""
    need = {header_key(n) for n in must_have}
    for i, row in enumerate(rows):
        if i >= scan:
            break
        hdr = [header_key(v) for v in (row or ())]
        if need <= set(hdr):
            return i, hdr
    return None, None


# --------------------------------------------------------------------------- #
# Cell values
# --------------------------------------------------------------------------- #
def text(v) -> Optional[str]:
    """Trimmed text; None for empty and Excel error cells. 'Blank' is kept (the payer
    rules decide what it means)."""
    if v is None:
        return None
    if isinstance(v, float):
        if v != v:
            return None
        if v.is_integer():
            v = int(v)
    s = str(v).strip()
    if not s or is_error_value(s):
        return None
    return s


def name(v) -> Optional[str]:
    """Like text(), but every blank form ('Blank', spaces, ...) is None — for doctor and
    manufacturer names, where HCG typed 'Blank' into empty cells."""
    return None if is_blank(v) else text(v)


def upper(v) -> Optional[str]:
    s = text(v)
    return s.upper() if s else None


def code(v) -> Optional[str]:
    return norm_code(v)


def invoice_prefix(v) -> Optional[str]:
    """'KABHK-...' -> 'KABHK'. The rest of the invoice number is dropped here."""
    s = text(v)
    if not s:
        return None
    return s.split("-", 1)[0].strip().upper() or None


def number(v) -> Optional[float]:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return None if f != f else f
    s = str(v).strip().replace(",", "")
    if not s or is_error_value(s):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def money(v) -> Optional[float]:
    """MRP columns carry float noise (1.590000033) on ~45% of rows: 2 decimals."""
    f = number(v)
    return None if f is None else round(f, 2)


def to_date(v) -> Optional[date]:
    """datetime, Excel serial (pyxlsb never converts) or a date string -> date."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        try:
            return (EXCEL_BASE + timedelta(days=float(v))).date()
        except (OverflowError, ValueError):
            return None
    s = str(v).strip()
    if not s or is_error_value(s):
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%Y",
                "%d-%b-%Y", "%d-%b-%y"):
        try:
            return datetime.strptime(s[:19] if fmt.startswith("%Y-%m-%d %H") else s.split(" ")[0],
                                     fmt).date()
        except ValueError:
            pass
    try:
        return (EXCEL_BASE + timedelta(days=float(s))).date()
    except (OverflowError, ValueError):
        return None
