# tests/test_formulary_pipeline.py — the whole formulary build on tiny made-up workbooks.
#
# HCG's real files carry patient data and never enter the repo, so this builds a miniature
# copy of their layout in tmp_path (hash-named top folder, a consolidated sales workbook with
# a Raw sheet, a pivot sheet and the WS sheet with HCG's lookup columns, an Item Master, a
# Payer Master with its marker row above the header) and runs the real build into tmp
# folders. It checks the numbers, that nothing line-level or patient-level is published, and
# the abort rules. Every directory is passed explicitly: nothing here touches app/data.
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from app.etl.formulary import config as C
from app.etl.formulary.build import build, main
from app.etl.formulary.discover import BuildError

DASH = "Formulary Compliance and Spend Analysis Dashboard data"
LINE_COLS = ["Type", "LOCATIONNAME", "SALESINVOICENO", "SALESDATE", "UHID", "IP_OPNUMBER",
             "BILLCATEGORY", "PATNAME", "TREATINGDOCID", "TREATINGDOCNAME", "MATERIALCODE",
             "MATERIALLINEITEM", "BRANDNAME", "QTY", "MRP", "TOTALMRP", "COSTPRICE",
             "TOTALCOSTPRICE", "PAYER1", "BATCHCODE", "EXPIRYDATE", "LOCATIONID", "GENERICNAME",
             "ISSUECODE", "MANUFACTURERNAME"]
HCG_COLS = ["FY 26-27 - CATEGORY",
            "payer as per updated Master sheet as per credit controle team-New",
            "payer+category as per updated Master sheet", "Onco/Non Onco"]
KR, DR = ("KABHK", "HCG KR, Bangalore"), ("KABHD", "HCG DR, Bangalore")

# Item Master: code -> (flag, category). Codes are written as floats, like HCG's file.
MASTER = {
    101001: ("FORMULARY", "R"),
    101002: ("OUT OF FORMULARY", None),
    101003: ("FORMULARY", "P"),
    101004: ("NON FORMULARY", None),
    201001: ("NON FORMULARY", None),
    201002: ("FORMULARY", "P-0PD"),          # typo, published as FORMULARYP-OPD
    201003: ("FORMULARY", "CONSUMABLE"),     # contrast media: group A
    201004: ("OUT OF FORMULARY", None),
    201005: ("FORMULARY", "R-OPD"),
}
HCG_CLASS = {"101001": "FORMULARYR", "101002": "OUT OF FORMULARY", "101003": "FORMULARYP",
             "101004": "NON FORMULARY", "201001": "NON FORMULARY", "201002": "FORMULARYP-0PD",
             "201003": "FORMULARYCONSUMABLE", "201004": "OUT OF FORMULARY",
             "201005": "FORMULARYR-OPD", "300001": "#N/A"}
SEGMENT = {"1": "A-Onco", "2": "B-Non Onco", "3": "C-Consumables"}

# Payer Master rows in sheet order: (CUSTOMERNAME, Type, Final after Discussion)
PAYERS = [
    ("ECHS", "GOV", "GOV"),
    ("Central Government Health Scheme (CGHS)", "GOV", "CASH(TPA)"),   # first row wins
    ("CENTRAL GOVERNMENT HEALTH SCHEME (CGHS)", "GOV", "GOV"),
    ("STAR HEALTH INSURANCE", "CASH(TPA)", "CASH(TPA)"),
]
OVERRIDES = """payer_name,bucket,from_period,to_period,note
SPECIAL STATE SCHEME,GOV,,,not in Payer Master; set by HCG
STAR HEALTH INSURANCE,GOV,2026-07,2026-07,month-scoped test row
"""
HOSPITAL_MAP = """his_code,his_location_name,his_location_id,sap_plant,sap_plant_name,confidence,note
KABHK,"HCG KR, Bangalore",10107,HC05,"HCG KR, Bangalore",high,test
KABHD,"HCG DR, Bangalore",10108,HC06,"HCG DR, Bangalore",high,test
"""


def L(hosp, typ, code, cost, payer=None, hcg_payer="CASH", doc=("D1", "Dr A"), qty=1, mrp=None,
      day=5):
    """One sales line: (hospital, type, material, cost, payer, HCG's payer, doctor)."""
    return dict(hosp=hosp, typ=typ, code=code, cost=cost, payer=payer, hcg_payer=hcg_payer,
                doc=doc, qty=qty, mrp=mrp if mrp is not None else (cost or 0) * 1.5, day=day)


LINES = [
    L(KR, "IP", 101001, 1000.0, "ECHS", "GOV"),                                    # A onco, R, GOV
    L(KR, "IP", 101002, 200.0, None, "CASH"),                                      # B onco
    L(KR, "OP", 201001, 300.0, "Blank ", "CASH", doc=("D2", "Dr B")),              # C non-onco
    L(KR, "IP", 201002, 400.0, "Special  State Scheme", "GOV", doc=("D2", "DR B")),  # A override
    L(KR, "IP", 201003, 50.0, "", "CASH"),                                         # A consumable class
    L(KR, "IP", 300001, 70.0, "", "CASH"),                                         # consumable
    L(KR, "IP", 101003, 500.0, "ECHS", "GOV", doc=(None, None)),                   # A onco P GOV
    L(KR, "IP", 201004, 100.0, "New Insurer Ltd", "CASH(TPA)", doc=("D2", "Dr B")),  # B, unmatched
    L(KR, "IP", 201005, None, "Central Government Health Scheme (CGHS)", "CASH(TPA)"),  # empty cost
    L(KR, "OP", 201001, 20.0, None, "CASH", doc=("D2", "Dr B")),                   # C
    L(DR, "IP", 101001, 800.0, "Star Health Insurance", "CASH(TPA)", doc=("D9", "Dr Z")),
    L(DR, "IP", 101002, 150.0, "Star Health Insurance", "CASH(TPA)", doc=("D9", "Dr Z")),
    L(DR, "OP", 201004, 60.0, None, "CASH", doc=("D9", "Dr Z")),
    L(DR, "OP", 101004, 40.0, None, "CASH", doc=("D9", "Dr Z")),
    L(DR, "IP", 300001, 90.0, "ECHS", "GOV", doc=("D9", "Dr Z")),
]
SECRET_UHIDS = {f"UH{i:06d}" for i in range(len(LINES))}


def _sales_rows(lines, period="2026-08", hcg=True):
    y, m = map(int, period.split("-"))
    out = []
    for i, ln in enumerate(lines):
        code = str(ln["code"])
        hosp_code, hosp_name = ln["hosp"]
        row = {
            "Type": ln["typ"], "LOCATIONNAME": hosp_name,
            "SALESINVOICENO": f"{hosp_code}-{ln['typ']}-{900000 + i}",
            "SALESDATE": datetime(y, m, ln["day"], 0, 0), "UHID": f"UH{i:06d}",
            "IP_OPNUMBER": f"IPN{i:05d}", "BILLCATEGORY": "CASH" if not ln["payer"] else "CREDIT",
            "PATNAME": f"Patient Name {i}", "TREATINGDOCID": ln["doc"][0],
            "TREATINGDOCNAME": ln["doc"][1], "MATERIALCODE": int(code),
            "MATERIALLINEITEM": f"ITEM {code}", "BRANDNAME": f"BRAND {code}", "QTY": ln["qty"],
            "MRP": ln["mrp"], "TOTALMRP": ln["mrp"], "COSTPRICE": ln["cost"],
            "TOTALCOSTPRICE": ln["cost"], "PAYER1": ln["payer"], "BATCHCODE": f"B{i}",
            "EXPIRYDATE": datetime(2028, 1, 1), "LOCATIONID": None, "GENERICNAME": f"GEN {code}",
            "ISSUECODE": f"PMIS{i}", "MANUFACTURERNAME": "ACME",
        }
        if hcg:
            row.update({HCG_COLS[0]: HCG_CLASS[code], HCG_COLS[1]: ln["hcg_payer"],
                        HCG_COLS[2]: "x", HCG_COLS[3]: SEGMENT[code[0]]})
        out.append(row)
    return out


def write_sales(path: Path, lines, period="2026-08", hcg=True, ws_name="IP&OP Con WS August 26"):
    wb = Workbook()
    raw = wb.active
    raw.title = "IP&OP Con Raw August 26"          # line sheet WITHOUT HCG's columns
    raw.append(LINE_COLS)
    for r in _sales_rows(lines, period, hcg=False):
        raw.append([r[c] for c in LINE_COLS])
    pivot = wb.create_sheet("Sheet1")
    pivot.append([None])
    pivot.append(["Sum of QTY", "Column Labels"])
    if hcg:
        ws = wb.create_sheet(ws_name)
        cols = LINE_COLS + HCG_COLS
        ws.append(cols)
        for r in _sales_rows(lines, period, hcg=True):
            ws.append([r[c] for c in cols])
    wb.save(path)


def write_master(path: Path, master=MASTER):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Material", "Material Description", "Minor Group", "Manufacturer",
               "Formulary flagging", "Category", "Materialtype"])
    for code, (flag, cat) in master.items():
        ws.append([float(code), f"MASTER DESC {code}", f"MINOR {code}", "ACME PHARMA", flag, cat,
                   "Onco" if str(code)[0] == "1" else "Non onco"])
    wb.save(path)


def write_payer_master(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([None, None, None, "CC", "CC", "ü"])                 # marker row
    ws.append(["LOCATIONID", "LOCATIONNAME", "CUSTOMERNAME", "Type",
               "Final after Discussion 5.2.2026", "AGREEMENTNAME"])
    for name, typ, final in PAYERS:
        ws.append([10107, "HCG KR, Bangalore", name, typ, final, f"{name} IP"])
    wb.save(path)


@pytest.fixture()
def env(tmp_path):
    root = tmp_path / "raw" / "1829e863"
    dash = root / DASH
    (dash / "IP and OP sales").mkdir(parents=True)
    (dash / "Item Master Sheet").mkdir()
    write_sales(dash / "IP and OP sales" / "IP & OP Consolidated File August 26.xlsx", LINES)
    write_master(dash / "Item Master Sheet" /
                 "Copy of Item master for formulary complience August -26_Updated.xlsx")
    write_payer_master(dash / "Copy of Copy of Payer Master Final Sheet.xlsx")
    ref = tmp_path / "reference"
    ref.mkdir()
    (ref / C.HOSPITAL_MAP_FILE).write_text(HOSPITAL_MAP)
    (ref / C.PAYER_OVERRIDES_FILE).write_text(OVERRIDES)
    dirs = dict(raw=tmp_path / "raw", cache=tmp_path / "cache", data=tmp_path / "data",
                reference=ref)
    return dict(dirs=dirs, dash=dash, tmp=tmp_path)


def run(env, **kw):
    return build(log=lambda *_: None, **env["dirs"], **kw)


def tables(env) -> dict:
    d = env["dirs"]["data"]
    return {t: pd.read_parquet(d / f"{t}.parquet") for t in C.TABLES}


def abc(df):
    d = df[df["category"].isin(["Onco Drugs", "Other Drugs"])]
    g = d.groupby("formulary_group")["cost"].sum()
    return tuple(round(float(g.get(x, 0.0)), 2) for x in "ABC")


# --- the numbers -------------------------------------------------------------------------------
def test_build_publishes_the_four_tables_and_they_reconcile(env):
    rep = run(env)
    t = tables(env)
    s = t[C.SUMMARY]
    assert int(s["lines"].sum()) == len(LINES)
    assert s["cost"].sum() == pytest.approx(sum(ln["cost"] or 0 for ln in LINES))
    assert s["mrp"].sum() == pytest.approx(sum(ln["mrp"] for ln in LINES))
    assert set(s["period"]) == {"2026-08"}
    assert set(s["plant"]) == {"HC05", "HC06"} and set(s["hospital"]) == {"KABHK", "KABHD"}
    assert set(s["hospital_name"]) == {KR[1], DR[1]}
    # drug-only tables reconcile with the summary's drug rows
    drugs = s[s["category"] != "Consumables"]
    for name in (C.BY_DOCTOR, C.BY_ITEM):
        assert t[name]["cost"].sum() == pytest.approx(drugs["cost"].sum())
        assert int(t[name]["lines"].sum()) == int(drugs["lines"].sum())
    bc = drugs[drugs["formulary_group"].isin(["B", "C"])]
    assert t[C.DOCTOR_ITEM]["cost"].sum() == pytest.approx(bc["cost"].sum())
    assert set(t[C.DOCTOR_ITEM]["formulary_group"]) == {"B", "C"}
    assert rep["months"]["2026-08"]["published"]["cost_diff"] == 0


def test_kr_formulary_split_and_compliance(env):
    rep = run(env)
    s = tables(env)[C.SUMMARY]
    kr = s[s["plant"] == "HC05"]
    # A = 1000 + 400 + 50 + 500 + 0 (empty cost), B = 200 + 100, C = 300 + 20
    assert abc(kr) == (1950.0, 300.0, 320.0)
    k = rep["months"]["2026-08"]["kr"]
    assert k["pooled_pct"] == pytest.approx(1950 / 2250 * 100, abs=1e-4)
    assert k["onco_pct"] == pytest.approx(1500 / 1700 * 100, abs=1e-4)
    assert k["non_onco_pct"] == pytest.approx(450 / 550 * 100, abs=1e-4)
    assert k["mean_pct"] == pytest.approx((1500 / 1700 + 450 / 550) / 2 * 100, abs=1e-4)
    # all hospitals: + DR A 800, B 150 + 60, C 40
    assert abc(s) == (2750.0, 510.0, 360.0)
    assert rep["months"]["2026-08"]["all_hospitals"]["pooled_pct"] == \
        pytest.approx(2750 / 3260 * 100, abs=1e-4)


def test_classes_categories_and_consumables(env):
    run(env)
    s = tables(env)[C.SUMMARY]
    cls = dict(zip(s["formulary_class"], s["formulary_group"]))
    assert cls["FORMULARYP-OPD"] == "A"                 # 'P-0PD' canonicalised
    assert cls["FORMULARYCONSUMABLE"] == "A"
    assert cls["FORMULARYR-OPD"] == "A"
    cons = s[s["category"] == "Consumables"]
    assert set(cons["formulary_class"]) == {"NOT IN MASTER"} and set(cons["formulary_group"]) == {"X"}
    assert cons["cost"].sum() == pytest.approx(160.0)
    item = tables(env)[C.BY_ITEM]
    assert "300001" not in set(item["material"])
    row = item[item["material"] == "101001"].iloc[0]
    assert row["material_desc"] == "MASTER DESC 101001" and row["generic_name"] == "MINOR 101001"
    assert row["category"] == "Onco Drugs"


def test_payer_buckets(env):
    run(env)
    s = tables(env)[C.SUMMARY]
    kr = s[s["plant"] == "HC05"].groupby("payer_bucket")["cost"].sum()
    # CASH: blank/'Blank '/'' payers (200 + 300 + 50 + 70 + 20); GOV: ECHS 1000 + 500 and the
    # override 400; CASH(TPA): unmatched 100 and CGHS (first master row) 0
    assert kr.to_dict() == pytest.approx({"CASH": 640.0, "GOV": 1900.0, "CASH(TPA)": 100.0})


def test_validation_report(env):
    rep = run(env)
    v = json.loads((env["dirs"]["data"] / C.VALIDATION_FILE).read_text())
    assert v["periods"] == ["2026-08"] and v["latest"] == "2026-08"
    m = v["months"]["2026-08"]
    assert m["source"]["sheet"] == "IP&OP Con WS August 26"      # not the Raw sheet
    assert m["source"]["rows_read"] == len(LINES) and m["source"]["has_hcg_columns"] is True
    assert m["item_master"]["materials"] == len(MASTER)
    assert m["drugs"]["found_in_master_pct"] == 100.0
    for what in ("category", "payer", "segment"):
        assert m["agreement"][what]["lines_pct"] == 100.0, (what, m["agreement"][what])
    assert m["agreement"]["category"]["by_hospital"]["KABHK"]["lines"] == 10
    assert m["payers"]["unmatched_top50"] == [{"name": "NEW INSURER LTD", "lines": 1, "cost": 100.0}]
    assert m["payers"]["by_source"]["override"]["lines"] == 1
    assert v["payer_master"]["conflicting_names"][0]["name"] == "CENTRAL GOVERNMENT HEALTH SCHEME (CGHS)"
    assert m["payers"]["conflicting_master_names_used"][0]["lines"] == 1
    assert m["kr"]["scheme"] == {"restricted": 1000.0, "preferred": 500.0}
    assert set(v["tables"]) == set(C.TABLES)
    assert rep["months"]["2026-08"]["period"] == "2026-08"


def test_nothing_line_level_or_patient_level_is_published(env):
    run(env)
    d = env["dirs"]["data"]
    blob = (d / C.VALIDATION_FILE).read_text()
    for name, df in tables(env).items():
        cols = {c.upper() for c in df.columns}
        assert not cols & {"UHID", "PATNAME", "IP_OPNUMBER", "SALESINVOICENO", "BATCHCODE",
                           "ISSUECODE", "EXPIRYDATE", "SALESDATE", "SALES_DATE"}, name
        assert not any(pd.api.types.is_datetime64_any_dtype(df[c]) for c in df.columns), name
        blob += df.astype(str).to_csv()
    for secret in list(SECRET_UHIDS) + ["Patient Name", "IPN0", "-IP-9000", "PMIS", "2026-08-05"]:
        assert secret not in blob, secret
    assert len(tables(env)[C.SUMMARY]) < len(LINES)          # aggregated, not one row per line


def test_doctor_names_and_unknown_doctor(env):
    run(env)
    doc = tables(env)[C.BY_DOCTOR]
    names = dict(zip(doc["doctor_id"], doc["doctor_name"]))
    assert names["D2"] == "Dr B"                    # 'Dr B' x3 beats 'DR B' x1
    assert names["UNKNOWN"] == "Not recorded"
    assert names["D9"] == "Dr Z"


def test_the_line_cache_stays_outside_the_data_dir_and_drops_patient_columns(env):
    run(env)
    cache = env["dirs"]["cache"]
    lines = pd.read_parquet(cache / "lines_2026-08.parquet")
    assert len(lines) == len(LINES)
    assert not {c.upper() for c in lines.columns} & {"UHID", "PATNAME", "IP_OPNUMBER",
                                                      "SALESINVOICENO", "BATCHCODE", "ISSUECODE"}
    assert set(lines["his_code"]) == {"KABHK", "KABHD"}
    assert not list(env["dirs"]["data"].glob("lines_*"))


# --- rebuilds ------------------------------------------------------------------------------
def _add_july(env, lines=None):
    dash = env["dash"]
    write_sales(dash / "IP and OP sales" / "IP and OP Sales Consolidated file July 2026.xlsx",
                lines or LINES[:6], period="2026-07", ws_name="IP&OP Jul 26 WS")
    write_master(dash / "Item Master Sheet" /
                 "Copy of item master for formulary complience July 2026.xlsx")


def test_rebuilding_one_month_replaces_only_that_month(env):
    _add_july(env)
    run(env)
    before = tables(env)[C.SUMMARY]
    jul_before = before[before["period"] == "2026-07"].reset_index(drop=True)
    assert set(before["period"]) == {"2026-07", "2026-08"}
    # a corrected August file arrives
    changed = [dict(ln) for ln in LINES]
    changed[0]["cost"] = 5000.0
    write_sales(env["dash"] / "IP and OP sales" / "IP & OP Consolidated File August 26.xlsx", changed)
    run(env, months=["2026-08"])
    after = tables(env)[C.SUMMARY]
    pd.testing.assert_frame_equal(after[after["period"] == "2026-07"].reset_index(drop=True),
                                  jul_before)
    aug = after[after["period"] == "2026-08"]
    assert aug["cost"].sum() == pytest.approx(sum(ln["cost"] or 0 for ln in changed))
    v = json.loads((env["dirs"]["data"] / C.VALIDATION_FILE).read_text())
    assert v["periods"] == ["2026-07", "2026-08"]


def test_overrides_are_month_scoped(env):
    _add_july(env, lines=[L(DR, "IP", 101001, 800.0, "Star Health Insurance", "GOV")])
    run(env)
    s = tables(env)[C.SUMMARY]
    star = s[(s["plant"] == "HC06") & (s["formulary_class"] == "FORMULARYR")]
    assert dict(zip(star["period"], star["payer_bucket"])) == {"2026-07": "GOV",
                                                                "2026-08": "CASH(TPA)"}


def test_from_cache_gives_the_same_tables(env):
    run(env)
    first = tables(env)
    run(env, from_cache=True)
    second = tables(env)
    for t in C.TABLES:
        pd.testing.assert_frame_equal(first[t], second[t])


def _nothing_published(env):
    d = env["dirs"]["data"]
    return not d.exists() or not (list(d.glob("*.parquet")) or list(d.glob("*.json")))


def test_dry_run_and_limit_never_publish(env):
    run(env, dry_run=True)
    assert _nothing_published(env)
    rep = run(env, limit=5)
    assert rep["months"]["2026-08"]["source"]["rows_read"] == 5
    assert _nothing_published(env)
    assert not (env["dirs"]["cache"] / "lines_2026-08.parquet").exists() or \
        not json.loads((env["dirs"]["cache"] / "lines_2026-08.json").read_text())["limited"]


# --- aborts ----------------------------------------------------------------------------------


def test_a_file_that_reads_zero_rows_aborts(env):
    write_sales(env["dash"] / "IP and OP sales" / "IP & OP Consolidated File August 26.xlsx", [])
    with pytest.raises(BuildError, match="0 rows"):
        run(env)
    assert _nothing_published(env)


def test_a_month_without_its_own_item_master_aborts(env):
    (env["dash"] / "Item Master Sheet" /
     "Copy of Item master for formulary complience August -26_Updated.xlsx").rename(
        env["dash"] / "Item Master Sheet" / "Copy of item master July 2026.xlsx")
    with pytest.raises(BuildError, match="no Item Master for 2026-08"):
        run(env)
    assert _nothing_published(env)


def test_an_unmapped_hospital_aborts(env):
    lines = LINES + [L(("MHNK1", "HCG Manavata Cancer Centre , Nashik Unit-1"), "IP", 101001, 5.0)]
    write_sales(env["dash"] / "IP and OP sales" / "IP & OP Consolidated File August 26.xlsx", lines)
    with pytest.raises(BuildError, match="MHNK1"):
        run(env)
    assert _nothing_published(env)


def test_a_sales_date_outside_the_file_month_aborts(env):
    write_sales(env["dash"] / "IP and OP sales" / "IP & OP Consolidated File August 26.xlsx",
                LINES, period="2026-07")
    with pytest.raises(BuildError, match="SALESDATE outside 2026-08"):
        run(env)
    assert _nothing_published(env)


def test_drugs_missing_from_the_master_abort(env):
    master = {k: v for k, v in MASTER.items() if k != 101001}
    write_master(env["dash"] / "Item Master Sheet" /
                 "Copy of Item master for formulary complience August -26_Updated.xlsx", master)
    with pytest.raises(BuildError, match="Item Master"):
        run(env)
    assert _nothing_published(env)


def test_overlapping_override_rows_are_rejected(env):
    ref = env["dirs"]["reference"] / C.PAYER_OVERRIDES_FILE
    ref.write_text(OVERRIDES + "Special State Scheme,CASH(TPA),2026-01,,overlaps the open row\n")
    with pytest.raises(BuildError, match="overlaps"):
        run(env)


def test_a_month_without_hcg_columns_builds_and_skips_the_agreement_check(env):
    write_sales(env["dash"] / "IP and OP sales" / "IP & OP Consolidated File August 26.xlsx",
                LINES, hcg=False)
    rep = run(env)
    m = rep["months"]["2026-08"]
    assert m["source"]["sheet"] == "IP&OP Con Raw August 26"
    assert m["agreement"] is None
    assert any("agreement" in w for w in m["warnings"])
    assert abc(tables(env)[C.SUMMARY]) == (2750.0, 510.0, 360.0)


def test_cli(env, monkeypatch, capsys):
    d = env["dirs"]
    for key, name in (("FORMULARY_RAW_DIR", "raw"), ("FORMULARY_CACHE_DIR", "cache"),
                      ("FORMULARY_DATA_DIR", "data"), ("REFERENCE_DIR", "reference")):
        monkeypatch.setenv(key, str(d[name]))
    assert main(["--list"]) == 0
    out = capsys.readouterr().out
    assert "2026-08 (Aug-26)" in out and "not cached" in out
    assert main(["--months", "2026-08"]) == 0
    assert (d["data"] / C.VALIDATION_FILE).exists()
    assert main(["--months", "2026-09"]) == 1
    assert "no sales data for 2026-09" in capsys.readouterr().err
    with pytest.raises(SystemExit):
        main(["--months", "Aug-26"])
