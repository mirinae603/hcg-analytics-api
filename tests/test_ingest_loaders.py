# tests/test_ingest_loaders.py — the legacy curated import (app/etl/ingest.py, run_etl,
# transforms) on tiny made-up extracts.
#
# HCG's SAP extracts never enter the repo, so each test writes a miniature copy of their layout
# in tmp_path with openpyxl: padded headers, dates as real dates and as 'dd.mm.yyyy' text, a
# consolidated consumption extract overlapping monthly ones, monthly extracts that book issues
# as negative quantities, GRN returns with positive amounts, deleted PO lines, Net Price per
# Price Unit, lines without a material, a capex tab, an empty Book.xlsx, and document numbers
# that must never reach a frame or a report. Every directory is passed explicitly: nothing here
# reads master/ or writes app/data.
from __future__ import annotations

import json
import os
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from app.core.config import Settings
from app.etl import ingest as I
from app.etl import run_etl, transforms
from app.etl.formulary.discover import BuildError

HASH = "1829e863"
CONS_DIR = "Consumption Dec 25 to May 26"
GRN_DIR = "GRN Dec 25 to May 2026"
INV_DIR = "Inventory report as on 31 May 2026"
PO_DIR = "PO details Dec 25 to May 26"
QUIET = lambda *_: None                                   # noqa: E731

# Document numbers sitting in the source columns: none may reach a frame or a report.
PO_A, PO_B, PO_C, PO_D, PO_E = "4100715422", "4100715423", "4100715424", "4100799001", "4100799002"
GR_NO, MAT_DOC, REF = "5000777001", "4900123456", "REF-SECRET-9"
SECRETS = (PO_A, PO_B, PO_C, PO_D, PO_E, GR_NO, MAT_DOC, REF)


def write_book(path: Path, header, rows, sheet="Sheet1", extra=()):
    """One workbook: `header` + rows (dicts keyed by header) on `sheet`, then any extra
    (name, header, rows) sheets."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for name, hdr, lines in ((sheet, header, rows), *extra):
        ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        ws.append(hdr)
        for ln in lines:
            ws.append([ln.get(h) for h in hdr])
    wb.save(path)


def assert_no_secrets(*things):
    text = " ".join(t.to_csv() if isinstance(t, pd.DataFrame) else json.dumps(t, default=str)
                    for t in things)
    for s in SECRETS:
        assert s not in text, f"document number {s} leaked"


# --------------------------------------------------------------------------- #
# Consumption
# --------------------------------------------------------------------------- #
CONS_HEADER = ["Material", "Material Description", "Plnt", "SLoc", "Reference", "Batch", "MvT",
               "Mat. Doc.", "Pstng Date", "Doc. Date", " Quantity", "EUn", "   Amount in LC",
               "Cost Ctr", "Customer", "MvtTypeTxt", "User", "Sales Val.", "Entry Date"]


def CR(material, plant, when, qty, amount, mvt=201):
    return {"Material": float(material), "Material Description": f"ITEM {material}",
            "Plnt": plant, "SLoc": "S124", "Reference": REF, "Batch": "B1", "MvT": mvt,
            "Mat. Doc.": float(MAT_DOC), "Pstng Date": when, "Doc. Date": when,
            " Quantity": qty, "EUn": "EA", "   Amount in LC": amount, "Cost Ctr": 1010101.0,
            "Customer": None, "MvtTypeTxt": "GI for cost center", "User": "USER1",
            "Sales Val.": 0, "Entry Date": when}


def SUM_LINE(amount):                       # SAP's total line at the bottom: no date, no movement
    return {"   Amount in LC": amount}


CONSOLIDATED = [CR(300001, "HC05", datetime(2025, 12, 3), 10, 100.0),
                CR(300002, "HT02", datetime(2025, 12, 20), 5, 50.0),     # a lab plant only it has
                CR(300001, "HC05", datetime(2026, 1, 7), 4, 40.0),
                CR(300003, "HC05", datetime(2026, 1, 8), 2, 30.0, mvt=202),  # not an issue
                SUM_LINE(220.0)]
MONTHLY = {
    "09.Consumption Report Dec 2025.xlsx": [CR(300001, "HC05", datetime(2025, 12, 3), 10, 100.0)],
    "10.Consumption Report Jan 2026.xlsx": [CR(300001, "HC05", datetime(2026, 1, 7), -4, -40.0)],
    # June books issues as NEGATIVE quantities, with text dates (4 June, not 6 April)
    "03.Consumption Report June 2026.xlsx": [CR(300001, "HC05", "04.06.2026", -3, -30.0),
                                             CR(300004, "HC06", "12.06.2026", -6, -66.0),
                                             CR(300004, "HC06", "13.06.2026", 0, 0.0),
                                             SUM_LINE(-96.0)],
    "04.Consumption Report-1 July 2026.xlsx": [CR(300001, "HC05", datetime(2026, 7, 2), 7, 70.0)],
}


def write_consumption(folder: Path, monthly=MONTHLY):
    folder.mkdir(parents=True, exist_ok=True)
    write_book(folder / "Consumption report Dec 25 to May 26.xlsx", CONS_HEADER, CONSOLIDATED)
    for name, rows in monthly.items():
        write_book(folder / name, CONS_HEADER, rows)


@pytest.fixture()
def raw(tmp_path):
    root = tmp_path / "master"
    write_consumption(root / HASH / CONS_DIR)
    return root


def test_each_month_comes_from_exactly_one_source(raw):
    got = I.load_consumption(raw, log=QUIET)
    by = {m["period"]: m for m in got.report["months"]}
    assert list(by) == ["2025-12", "2026-01", "2026-06", "2026-07"]
    assert by["2025-12"]["file"] == by["2026-01"]["file"] == "Consumption report Dec 25 to May 26.xlsx"
    assert by["2025-12"]["source"] == "consolidated"
    assert by["2026-06"]["file"] == "03.Consumption Report June 2026.xlsx"
    assert by["2026-07"]["source"] == "monthly"
    superseded = {s["file"] for s in got.report["superseded"]}
    assert superseded == {"09.Consumption Report Dec 2025.xlsx", "10.Consumption Report Jan 2026.xlsx"}
    per = got.frame.groupby(got.frame["posting_date"].dt.strftime("%Y-%m")).agg(
        lines=("qty", "size"), qty=("qty", "sum"), amount=("amount_lc", "sum"))
    assert per.loc["2025-12"].tolist() == [2, 15, 150]          # the monthly Dec copy is not added
    assert per.loc["2026-01"].tolist() == [1, 4, 40]            # the 202 line is not an issue
    assert per.loc["2026-06"].tolist() == [2, 9, 96]            # signs flipped, zero line dropped
    assert per.loc["2026-07"].tolist() == [1, 7, 70]


def test_issues_are_positive_and_text_dates_are_day_first(raw):
    got = I.load_consumption(raw, log=QUIET)
    f = got.frame
    assert (f["qty"] > 0).all() and (f["amount_lc"] > 0).all()
    june = f[f["posting_date"].dt.month == 6].sort_values("posting_date")
    assert [d.date() for d in june["posting_date"]] == [date(2026, 6, 4), date(2026, 6, 12)]
    by = {m["period"]: m for m in got.report["months"]}
    assert by["2026-06"]["issues_booked"].startswith("negative")
    assert by["2026-07"]["issues_booked"] == "positive"
    assert by["2026-06"]["no_quantity"] == 1


def test_consumption_frame_keeps_the_fact_schema_and_no_document_number(raw):
    got = I.load_consumption(raw, log=QUIET)
    f = got.frame
    assert list(f.columns) == I.FACT_CONSUMPTION_COLUMNS
    assert (f["mvt"] == 201).all()
    assert set(f["material"]) == {"300001", "300002", "300004"}         # codes, never '300001.0'
    assert set(f["cost_ctr"]) == {"1010101"}
    assert f["month"].tolist().count("June") == 2 and set(f["year"]) == {2025, 2026}
    assert_no_secrets(f, got.report)


def test_two_monthly_files_for_one_month_abort(tmp_path):
    folder = tmp_path / "master" / HASH / CONS_DIR
    write_consumption(folder, {**MONTHLY, "03.Consumption Report June 2026 (copy).xlsx":
                               MONTHLY["03.Consumption Report June 2026.xlsx"]})
    with pytest.raises(BuildError, match="2026-06"):
        I.load_consumption(tmp_path / "master", log=QUIET)


def test_two_consolidated_files_for_one_month_abort(tmp_path):
    folder = tmp_path / "master" / HASH / CONS_DIR
    write_consumption(folder)
    write_book(folder / "Consumption report Jan to Feb 26.xlsx", CONS_HEADER,
               [CR(300001, "HC05", datetime(2026, 1, 9), 1, 10.0),
                CR(300001, "HC05", datetime(2026, 2, 9), 1, 10.0)])
    with pytest.raises(BuildError, match="2026-01"):
        I.load_consumption(tmp_path / "master", log=QUIET)


def test_a_source_booking_issues_both_ways_aborts(tmp_path):
    folder = tmp_path / "master" / HASH / CONS_DIR
    write_consumption(folder, {"03.Consumption Report June 2026.xlsx":
                               [CR(300001, "HC05", "04.06.2026", -3, -30.0),
                                CR(300001, "HC05", "05.06.2026", 3, 30.0)]})
    with pytest.raises(BuildError, match="one way"):
        I.load_consumption(tmp_path / "master", log=QUIET)


def test_an_empty_workbook_is_skipped_and_reported(raw):
    Workbook().save(raw / HASH / CONS_DIR / "Book.xlsx")
    got = I.load_consumption(raw, log=QUIET)
    assert got.report["skipped"] == [{"file": "Book.xlsx", "reason": "empty workbook"}]


# --------------------------------------------------------------------------- #
# GRN (through the purchase build's reader)
# --------------------------------------------------------------------------- #
GRN_HEADER = ["Sl.No", "PO No", "PO Date", "GR No", "mvt type", "reference document", "GR Date",
              "PO to GR TAT", "PR to GR TAT", "Goods supplier", "Goods supplier name",
              "Vendor Inv No", "Plant", "S.Loc", "Material", "Material Description",
              "Expiry Date", "GR Qty", "GR UOM", "    Base Price", "Discount Value",
              "     Net Price", "Total Amount without tax", " Tax %", "   Tax Amount",
              "Total Amount with Tax", "      Unit MRP", "Total MRP Value w/o free",
              "Manufacturer", "Material Type", "Minor Category", "Major Group", "Minor Group",
              "Formulary"]


def GR(mvt, when, material, total, qty=1.0):
    return {"Sl.No": 1, "PO No": float(PO_A), "PO Date": "01.06.2026", "GR No": float(GR_NO),
            "mvt type": mvt, "reference document": REF, "GR Date": when, "PO to GR TAT": 3,
            "PR to GR TAT": None, "Goods supplier": 300100.0, "Goods supplier name": "MEDS",
            "Vendor Inv No": REF, "Plant": "HC05", "S.Loc": "S124",
            "Material": float(material) if material else None,
            "Material Description": f"ITEM {material}", "Expiry Date": datetime(2028, 1, 31),
            "GR Qty": qty, "GR UOM": "EA", "    Base Price": total / qty, "Discount Value": 0.0,
            "     Net Price": total / qty, "Total Amount without tax": total, " Tax %": 5.0,
            "   Tax Amount": total * 0.05, "Total Amount with Tax": total * 1.05,
            "      Unit MRP": 0.0, "Total MRP Value w/o free": 0.0, "Manufacturer": "ACME",
            "Material Type": "DRUG", "Minor Category": "DRUG", "Major Group": "DRUGS",
            "Minor Group": "MOLECULE", "Formulary": "FORMULARY"}


GRN_JUNE = [GR(101, "04.06.2026", 100001, 200.0, qty=2.0),        # text date: 4 June
            GR(101, datetime(2026, 6, 5), 300001, 30.0),
            GR(122, "05.06.2026", 100001, 80.0),                   # a return, amount POSITIVE
            GR(161, "06.06.2026", 200002, 20.0),                   # a return against the PO
            GR(101, "12.06.2026", None, 999.0)]                    # civil works: no material
GRN_JULY = [GR(101, datetime(2026, 7, 9), 100001, 70.0)]


def write_grn(folder: Path):
    folder.mkdir(parents=True, exist_ok=True)
    write_book(folder / "03.Consolidated GRN June 2026.xlsx", GRN_HEADER, GRN_JUNE, "GRN")
    write_book(folder / "04.Consolidated GRN July 2026.xlsx", GRN_HEADER, GRN_JULY, "GRN")


def test_grn_keeps_receipts_only_and_counts_what_it_leaves_out(tmp_path):
    write_grn(tmp_path / "master" / HASH / GRN_DIR)
    got = I.load_grn(tmp_path / "master", cache=tmp_path / "cache", log=QUIET)
    f = got.frame
    assert len(f) == 3 and f["total_amount_wo_tax"].sum() == 300.0
    june = got.report["files"][0]
    assert june["period"] == "2026-06" and june["lines"] == 2 and june["value"] == 230.0
    assert june["not_receipts"] == {"122": {"lines": 1, "value": 80.0},
                                    "161": {"lines": 1, "value": 20.0}}
    assert june["no_material"] == {"lines": 1, "value": 999.0}


def test_grn_keeps_the_fact_schema_with_day_first_dates_and_no_document_number(tmp_path):
    write_grn(tmp_path / "master" / HASH / GRN_DIR)
    got = I.load_grn(tmp_path / "master", cache=tmp_path / "cache", log=QUIET)
    f = got.frame
    assert list(f.columns) == I.FACT_GRN_COLUMNS
    assert not {"po_no", "gr_no"} & set(f.columns)                     # no document number column
    assert date(2026, 6, 4) in {d.date() for d in f["gr_date"]}
    assert (f["posting_date"] == f["gr_date"]).all()
    assert sorted(zip(f["year"], f["month"])) == [(2026, "July"), (2026, "June"), (2026, "June")]
    assert f["material"].tolist().count("100001") == 2
    assert_no_secrets(f, got.report)


# --------------------------------------------------------------------------- #
# PO
# --------------------------------------------------------------------------- #
PO_HEADER = ["Sl.No", "PO Date", "PO Doc.Type", "Doc.Type description", "PO No", "PO Li.No",
             "Del Flag", "PO Vendor", "Vendor Name", "Goods Supplier", "HSN Code",
             "Goods Supplier Name", "Material", "Material Description", "PO Quantity", "PO UOM",
             "Open Qty", "Delivery Date", "Net Price", "Currency", "Price Unit",
             "Order Price Unit", "Total value without Tax", "Plant", "Plant Name", "GL code",
             "Storage Location", "Tax Code", "Tax Code Description", "CGST %", "    CGST Value",
             "SGST %", "    SGST Value", "IGST %", "   IGST Value", "Additional value",
             "Total value with Tax", "Manufacturer", "Material Type", "Minor Category",
             "Major Group", "Minor Group", "Created by"]


def PR(po, when, material, qty, net, total, price_unit=1, flag=None, doc="Domestic PO",
       delivery=None):
    return {"Sl.No": 1, "PO Date": when, "PO Doc.Type": "ZP01", "Doc.Type description": doc,
            "PO No": float(po), "PO Li.No": 10, "Del Flag": flag, "PO Vendor": 10090.0,
            "Vendor Name": "MEDS", "Material": float(material) if material else None,
            "Material Description": f"ITEM {material}", "PO Quantity": qty, "PO UOM": "EA",
            "Open Qty": 0, "Delivery Date": delivery or when, "Net Price": net,
            "Currency": "INR", "Price Unit": price_unit, "Order Price Unit": "EA",
            "Total value without Tax": total, "Plant": "HC05", "Plant Name": "HCG KR",
            "Storage Location": "S124", "    CGST Value": total * 0.025,
            "    SGST Value": total * 0.025, "   IGST Value": 0.0,
            "Total value with Tax": total * 1.05, "Major Group": "DRUGS", "Minor Group": "MOL"}


PO_DEC = [PR(PO_A, datetime(2025, 12, 3), 200001, 10, 45.0, 450.0),
          PR(PO_A, datetime(2025, 12, 3), 200002, 20, 450.0, 900.0, price_unit=10),
          PR(PO_B, datetime(2025, 12, 4), 200003, 200003, 100.0, 20000300.0, flag="L"),
          PR(PO_B, datetime(2025, 12, 4), 200004, 1, 0.0, 0.0, flag="S"),   # blocked: still an order
          PR(PO_B, datetime(2025, 12, 5), None, 1, 5000.0, 5000.0, doc="Service PO")]
CAPEX_TAB = [PR(PO_C, datetime(2025, 12, 6), 600001, 1, 777.0, 777.0, doc="Dom Capital PO")]
PO_JUNE = [PR(PO_D, "04.06.2026", 200001, 5, 45.0, 225.0, delivery="10.06.2026"),
           PR(PO_E, "05.06.2026", 200005, -3, -40.0, -120.0)]           # a return, not an order


def write_po(folder: Path):
    folder.mkdir(parents=True, exist_ok=True)
    write_book(folder / "01.PO details for Dec 2025 extr 02.01.2026.xlsx", PO_HEADER, PO_DEC,
               "PO details for Dec 2025 (2)", extra=[("Capex", PO_HEADER, CAPEX_TAB)])
    write_book(folder / "03. PO details June 2026 extr 01.07.26.xlsx", PO_HEADER, PO_JUNE,
               "03. PO details June 2026 extr 0")
    Workbook().save(folder / "Book.xlsx")                         # HCG's empty workbook


@pytest.fixture()
def po_raw(tmp_path):
    write_po(tmp_path / "master" / HASH / PO_DIR)
    return tmp_path / "master"


def test_po_drops_deleted_lines_lines_without_a_material_and_returns(po_raw):
    got = I.load_po(po_raw, log=QUIET)
    f = got.frame
    assert sorted(f["material"]) == ["200001", "200001", "200002", "200004"]
    assert f["total_value_wo_tax"].sum() == 450.0 + 900.0 + 0.0 + 225.0
    assert (f["po_qty"] >= 0).all()
    dec, june = got.report["files"]
    assert dec["deleted"] == {"lines": 1, "value": 20000300.0}
    assert dec["no_material"] == {"lines": 1, "value": 5000.0, "by_doc_type": {"Service PO": 5000.0}}
    assert dec["qty_equals_material_code"] == 0          # the keying slip was a deleted line
    assert (dec["returns"], june["returns"]) == ({"lines": 0, "value": 0.0}, {"lines": 1, "value": -120.0})
    assert june["lines"] == 1 and june["lines_read"] == 2


def test_po_net_price_is_per_unit(po_raw):
    f = I.load_po(po_raw, log=QUIET).frame
    assert f.loc[f["material"] == "200002", "net_price"].tolist() == [45.0]
    assert ((f["po_qty"] * f["net_price"]).round(6) == f["total_value_wo_tax"]).all()


def test_po_reads_the_first_sheet_only_and_skips_the_empty_book(po_raw):
    got = I.load_po(po_raw, log=QUIET)
    assert "600001" not in set(got.frame["material"])               # the capex tab's line
    assert got.report["files"][0]["sheets_not_read"] == ["Capex"]
    assert got.report["skipped"] == [{"file": "Book.xlsx", "reason": "empty workbook"}]


def test_po_numbers_become_opaque_per_order_keys(po_raw):
    got = I.load_po(po_raw, log=QUIET)
    f = got.frame
    assert f["po_no"].str.fullmatch(r"PO-\d{7}").all()
    by = f.groupby("po_no")["material"].apply(sorted).to_dict()
    # PO_A's two lines share a key, PO_B's surviving line and PO_D's line get their own
    assert sorted(by.values()) == [["200001"], ["200001", "200002"], ["200004"]]
    assert f["po_no"].nunique() == 3
    assert_no_secrets(f, got.report)


def test_po_keeps_the_fact_schema_and_day_first_dates(po_raw):
    f = I.load_po(po_raw, log=QUIET).frame
    assert list(f.columns) == I.FACT_PO_COLUMNS
    june = f[f["month"] == "June"]
    assert june["po_date"].dt.date.tolist() == [date(2026, 6, 4)]
    assert june["delivery_date"].dt.date.tolist() == [date(2026, 6, 10)]
    assert sorted(set(zip(f["year"], f["month_num"]))) == [(2025, 12), (2026, 6)]


def test_a_po_workbook_without_the_po_header_aborts(po_raw):
    write_book(po_raw / HASH / PO_DIR / "07. PO details Oct 2026.xlsx",
               ["Purchasing Doc", "Item", "Material"], [{"Purchasing Doc": 1, "Material": 2}])
    with pytest.raises(BuildError, match="no PO header"):
        I.load_po(po_raw, log=QUIET)


def test_two_po_workbooks_for_one_month_abort(po_raw):
    write_book(po_raw / HASH / PO_DIR / "03. PO details June 2026 extr 05.07.26.xlsx", PO_HEADER,
               PO_JUNE)
    with pytest.raises(BuildError, match="2026-06"):
        I.load_po(po_raw, log=QUIET)


# --------------------------------------------------------------------------- #
# Inventory, folders and settings
# --------------------------------------------------------------------------- #
INV_HEADER = ["Date", "Plant", "Storage Location", "Storage Location Description", "Material",
              "Material Desc", "Batch", "Expiry Date", "GRN Date", "    Quantity", "Basic UOM",
              "   Total Cost", "    Total MRP", " Aging", "Material Group"]


def INV(material, qty, cost, grn, expiry, aging=10):
    return {"Date": "31.05.2026", "Plant": "HC05", "Storage Location": "S124",
            "Storage Location Description": "IP PHARMACY", "Material": float(material),
            "Material Desc": f"ITEM {material}", "Batch": "B1", "Expiry Date": expiry,
            "GRN Date": grn, "    Quantity": qty, "Basic UOM": "EA", "   Total Cost": cost,
            "    Total MRP": cost * 1.2, " Aging": aging, "Material Group": "M001-TABLETS"}


def write_inventory(folder: Path):
    folder.mkdir(parents=True, exist_ok=True)
    write_book(folder / "Consolidated Inventory report 02.06.2026.xlsx", INV_HEADER,
               [INV(300001, 100, 1000.0, "05.03.2026", datetime(2027, 1, 31)),
                INV(300004, 10, 50.0, datetime(2026, 5, 2), "12.01.2027", aging=200)])


def test_inventory_dates_are_day_first(tmp_path):
    write_inventory(tmp_path / "master" / HASH / INV_DIR)
    got = I.load_inventory(tmp_path / "master", log=QUIET)
    f = got.frame
    assert f["grn_date"].dt.date.tolist() == [date(2026, 3, 5), date(2026, 5, 2)]
    assert f["expiry_date"].dt.date.tolist() == [date(2027, 1, 31), date(2027, 1, 12)]
    assert got.report["snapshot"] == "2026-05-31"


def test_a_second_inventory_snapshot_aborts(tmp_path):
    folder = tmp_path / "master" / HASH / INV_DIR
    write_inventory(folder)
    write_book(folder / "Consolidated Inventory report 01.07.2026.xlsx", INV_HEADER, [])
    with pytest.raises(BuildError, match="one inventory snapshot"):
        I.load_inventory(tmp_path / "master", log=QUIET)


def test_source_folders_are_found_below_the_hashed_folder(tmp_path):
    (tmp_path / "master" / HASH / PO_DIR).mkdir(parents=True)
    assert I.find_source_dir(I.PO_DIR_RE, "PO", tmp_path / "master").name == PO_DIR
    with pytest.raises(BuildError, match="no 'Consumption"):
        I.find_source_dir(I.CONSUMPTION_DIR_RE, "'Consumption ...'", tmp_path / "master")
    (tmp_path / "master" / "other" / "PO details Jun 26").mkdir(parents=True)
    with pytest.raises(BuildError, match="more than one"):
        I.find_source_dir(I.PO_DIR_RE, "PO", tmp_path / "master")
    with pytest.raises(BuildError, match="does not exist"):
        I.find_source_dir(I.PO_DIR_RE, "PO", tmp_path / "nowhere")


def test_raw_data_dir_defaults_to_master_and_the_environment_wins(monkeypatch, tmp_path):
    monkeypatch.delenv("RAW_DATA_DIR", raising=False)
    root = Path(__file__).resolve().parents[1]
    assert Path(Settings().RAW_DATA_DIR) == root / "master"
    monkeypatch.setenv("RAW_DATA_DIR", str(tmp_path))
    assert Settings().RAW_DATA_DIR == str(tmp_path)


# --------------------------------------------------------------------------- #
# run_etl: the whole curated build on the miniature tree
# --------------------------------------------------------------------------- #
def test_build_curated_writes_every_table_and_the_source_report(raw, tmp_path):
    write_grn(raw / HASH / GRN_DIR)
    write_po(raw / HASH / PO_DIR)
    write_inventory(raw / HASH / INV_DIR)
    out = tmp_path / "curated"
    run_etl.build_curated(raw, cache=tmp_path / "cache", out=out)
    for name in ("fact_consumption", "fact_inventory", "fact_grn", "fact_po", "dim_plant",
                 "dim_material", "dim_sloc", "dim_vendor", "dim_costcenter"):
        assert (out / f"{name}.parquet").exists(), name
    rep = json.loads((out / "_dq_report.json").read_text())
    assert rep["periods"] == {"consumption": ["2025-12", "2026-01", "2026-06", "2026-07"],
                              "grn": ["2026-06", "2026-07"], "po": ["2025-12", "2026-06"]}
    assert set(rep["sources"]) == {"consumption", "inventory", "grn", "po"}
    assert rep["grn_rows"] == 3 and rep["po_rows"] == 4
    assert_no_secrets(rep, pd.read_parquet(out / "fact_po.parquet"),
                      pd.read_parquet(out / "fact_grn.parquet"))


# --------------------------------------------------------------------------- #
# transforms: what the new months change downstream
# --------------------------------------------------------------------------- #
@pytest.fixture()
def kpi_dir(tmp_path, monkeypatch):
    d = tmp_path / "kpi"
    monkeypatch.setattr(transforms, "KPI", d)
    return d


def _frame(rows, cols):
    return pd.DataFrame(rows, columns=cols)


DIM = _frame([("300001", "ITEM 300001", "M001"), ("300004", "ITEM 300004", "M001")],
             ["material", "material_desc", "material_group"])


def _cons(rows):
    f = _frame(rows, ["plant", "material", "posting_date", "qty", "amount_lc"])
    f["posting_date"] = pd.to_datetime(f["posting_date"])
    f["year"] = f["posting_date"].dt.year
    f["month"] = f["posting_date"].dt.month_name()
    f["cost_ctr"] = "1010101"
    return f


def test_stock_change_counts_only_months_both_sources_cover(kpi_dir):
    grn = _frame([("HC05", "300001", 2026, "July", 10.0), ("HC05", "300001", 2026, "August", 5.0)],
                 ["plant", "material", "year", "month", "gr_qty"])
    cons = _cons([("HC05", "300001", "2026-07-02", 4.0, 40.0),
                  ("HC05", "300001", "2026-09-02", 6.0, 60.0)])
    transforms.build_stock_change(grn, cons, DIM)
    out = pd.read_parquet(kpi_dir / "kpi_stock_change.parquet")
    assert out[["month", "inflow", "outflow", "stock_change"]].values.tolist() == [
        ["July", 10.0, 4.0, 6.0]]


def test_cycle_time_counts_receipt_lines_without_a_gr_number(kpi_dir):
    nan = float("nan")
    grn = _frame([("HC05", 2026, "June", "MEDS", 3.0, nan, 1.0),
                  ("HC05", 2026, "June", "MEDS", 5.0, nan, 1.0),
                  ("HC05", 2026, "June", "MEDS", 4.0, nan, 2.0)],
                 ["plant", "year", "month", "vendor_name", "po_to_gr_tat", "pr_to_gr_tat", "gr_qty"])
    po = _frame([("HC05", 10.0, 0.0)], ["plant", "po_qty", "open_qty"])
    transforms.build_additional_kpis(grn, po)
    assert pd.read_parquet(kpi_dir / "kpi_cycle_time.parquet")["gr_lines"].tolist() == [3]
    assert pd.read_parquet(kpi_dir / "kpi_vendor_lead_time.parquet")["gr_lines"].tolist() == [3]


def test_stock_kpis_use_the_consumption_up_to_the_snapshot(kpi_dir):
    snap = pd.Timestamp("2026-05-31")
    inv = _frame([("HC05", "300001", "ITEM 300001", "M001", 100.0, 1000.0, 1200.0, 10,
                   pd.Timestamp("2026-03-05"), pd.Timestamp("2027-01-31"), snap, "B1"),
                  ("HC05", "300004", "ITEM 300004", "M001", 10.0, 50.0, 60.0, 10,
                   pd.Timestamp("2026-05-02"), pd.Timestamp("2027-01-12"), snap, "B2")],
                 ["plant", "material", "material_desc", "material_group", "qty", "total_cost",
                  "total_mrp", "aging_days", "grn_date", "expiry_date", "snapshot_date", "batch"])
    cons = _cons([("HC05", "300001", "2025-12-01", 30.0, 300.0),
                  ("HC05", "300001", "2026-05-31", 30.0, 300.0),
                  ("HC05", "300004", "2026-07-02", 99.0, 990.0)])    # after the stock count
    assert len(transforms.consumption_until(cons, snap)) == 2
    transforms.build_inventory_kpis(inv, cons, DIM)
    doh = pd.read_parquet(kpi_dir / "kpi_doh.parquet").set_index("material")
    assert doh.loc["300001", "consumption_qty"] == 60.0
    assert doh.loc["300004", "consumption_qty"] == 0.0
    aging = pd.read_parquet(kpi_dir / "kpi_inventory_aging.parquet").set_index("material")
    assert (aging["age_since_last_sale_days"].dropna() >= 0).all()
    moving = pd.read_parquet(kpi_dir / "kpi_non_moving.parquet")
    assert moving["material"].tolist() == ["300004"]                # July's issue does not count
