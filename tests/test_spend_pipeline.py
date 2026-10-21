# tests/test_spend_pipeline.py — the spend build (Attis, scheme, cash & TPA) on tiny made-up files.
#
# Builds a miniature copy of HCG's folder (a consolidated sales workbook, an Item Master, the
# Payer Master and a Spend Analysis Main Sheet with HCG's SCHEME and CASH&TPA layouts) in
# tmp_path, plus a bifurcation reference file, runs the formulary build to fill a tmp line
# cache, then the spend build into a tmp folder. Every expected number is worked out by hand
# from the lines below. Nothing touches app/data.
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from app.etl.formulary import config as FC
from app.etl.formulary.build import build as build_formulary
from app.etl.formulary.discover import BuildError
from app.etl.spend import config as C
from app.etl.spend.build import build, main
from app.etl.spend.masters import load_bifurcation, load_cash_list, load_scheme_list
from tests.test_formulary_pipeline import (DASH, HOSPITAL_MAP, LINE_COLS, MASTER, OVERRIDES,
                                           write_master, write_payer_master)

KR, DR = ("KABHK", "HCG KR, Bangalore"), ("KABHD", "HCG DR, Bangalore")
ATTIS_MAKER = "ATTIS GLOBAL PHARMA"
ITEM_MASTER = {**MASTER, 101901: ("FORMULARY", "SCH"), 300001: ("FORMULARY", "CONSUMABLE")}

# HCG's SCHEME header, column for column (column 15 is HCG's empty spacer).
SCHEME_HEADER = [
    "Type", "CATEGORY", "COMPANY NAME", "SAP CODE", "SAP NAME", "Check (Vlook up)",
    "2026-27 HCG Pharma Selling  Price(Rs) Excl of GST", "2026-27 MRP(Rs) Incl of GST",
    "2026-27 GST %", "2026-27 Category", "2026-27 SHARING", "Abb MRP", "Unit Margin Value",
    "Unit Margin %", None, "SAP code", "SAP NAME", "Attis Brand/Other Brands",
    "HCG Pharma Selling Price(Rs) Excl of VAT", "Government Scheme",
    "used brand Total Cost value", "suggested brand Total value", "Margin loss"]
# used code -> (used price, FY category, suggested code, suggested price, Attis?)
SCHEME = [
    (101001, 1000.0, "R", 101901, 400.0, True),        # onco -> Attis brand
    (201002, 50.0, "P-0PD", 201902, 30.0, False),     # non-onco -> other brand (typo kept)
    (201002, 50.0, "P-0PD", 201902, 30.0, False),     # identical duplicate row: one switch
    (101003, 700.0, "P", None, None, None),           # no suggestion: not a switch
    (300001, 7.0, None, 300901, 5.0, False),          # consumable -> other brand
    (201004, 30.0, "P", 201904, None, False),         # suggestion without a price: skipped
    (101901, 400.0, "SCH", None, None, None),         # the Attis brand itself
]
SCHEME_GROUP = {300001: "SUTURES"}                     # -> HCG's J&J & Covidien vs Healthium line

# HCG's CASH&TPA header, column for column (column 14 is HCG's empty spacer).
CASH_HEADER = [
    "Type ", "COMPANY NAME", "SAP CODE", "SAP NAME", "2026-27 HCG Pharma Selling  Price(Rs) Excl of GST",
    "2026-27 MRP(Rs) Incl of GST", "2026-27 GST %", "2026-27 Category", "2026-27 SHARING",
    "Cash TPA Used Qty ( Vlookup from IP,OP Sales)", "Abb MRP", "Unit Margin Value", "Unit Margin %",
    None, "SAP code", "SAP NAME", "HCG Pharma Selling Price(Rs) Excl of VAT", "MRP(Rs) Incl of GST",
    "GST %", "Abb MRP", "Unit Margin Value", "Unit Margin %", "CASH & TPA Used Quantity",
    "used brand Total Cost value", "used brand Total MRP", "suggested brand Total value",
    "suggested brand Total MRP", "Actual", "Potential", "Margin loss", "Diffrence"]
# used code -> (price, MRP incl GST, GST, category, SHARING, suggested, its price, MRP, GST)
# Unit margin = MRP / 1.05 - price: 101001 1200 - 1000 = 200 vs 101950 1500 - 900 = 600 (+400);
# 201002 100 - 50 = 50 vs 150 - 30 = 120 (+70); 300001 20 - 7 = 13 vs 30 - 2 = 28 (+15).
CASH = [
    (101001, 1000.0, 1260.0, 5, "R", 0.1, 101950, 900.0, 1575.0, 5),
    (201002, 50.0, 105.0, 5, "P", 0.3, 201902, 30.0, 157.5, 5),
    (300001, 7.0, 21.0, 5, "P", 0.2, 300950, 2.0, 31.5, 5),
    (201004, 30.0, 63.0, None, "P", 0.1, 201904, 20.0, 84.0, 5),      # no GST: skipped
    (101003, 700.0, 840.0, 5, "P", 0.3, None, None, None, None),       # no suggestion
]
# 101001 is bifurcated; its molecule group also holds 101003 and the suggested 101950.
BIFURCATION = """material,flag,group,material_desc,note
101001,N,101001 101003 101950,LIST NAME 101001,test
"""


def S(hosp, typ, code, qty, cost, payer, doc=("D1", "Dr A"), maker="ACME"):
    return dict(hosp=hosp, typ=typ, code=code, qty=qty, cost=cost, payer=payer, doc=doc,
                maker=maker)


LINES = [
    S(KR, "IP", 101001, 2, 1800.0, "ECHS"),                                   # 1 Attis switch
    S(KR, "OP", 101001, 1, 900.0, "Special  State Scheme", doc=("D2", "Dr B")),  # 2 GOV by override
    S(KR, "IP", 101001, 5, 4500.0, None),                                     # 3 CASH: not counted
    S(KR, "IP", 201002, 10, 450.0, "ECHS", doc=("D2", "Dr B")),               # 4 other switch
    S(KR, "IP", 300001, 4, 30.0, "ECHS"),                                     # 5 consumable switch
    S(KR, "IP", 101003, 1, 500.0, "ECHS"),                                    # 6 no suggestion
    S(KR, "IP", 201004, 3, 90.0, "ECHS"),                                     # 7 no suggested price
    S(DR, "IP", 101001, 1, 950.0, "ECHS", doc=("D9", "Dr Z")),                # 8 Attis switch, DR
    S(KR, "IP", 101901, 2, 700.0, "ECHS", maker=ATTIS_MAKER),                 # 9 Attis product
    S(DR, "OP", 101901, 1, 350.0, None, doc=("D9", "Dr Z"), maker=ATTIS_MAKER),  # 10 Attis, CASH
    S(KR, "IP", 101001, -1, -900.0, "ECHS"),                                  # 11 a return
    S(KR, "OP", 101001, 3, 2700.0, "Star Health Insurance", doc=("D2", "Dr B")),  # 12 TPA
    S(KR, "IP", 201002, 6, 270.0, None),                                      # 13 cash
    S(KR, "IP", 300001, 10, 70.0, None),                                      # 14 cash
    S(DR, "IP", 101001, 2, 1900.0, None, doc=("D9", "Dr Z")),                 # 15 cash, DR
]
# By hand. Attis: used 2x1000 + 1000 + 1000 - 1000 = 3000, suggested 800 + 400 + 400 - 400
# = 1200. Other switches: 10 x 50 = 500 vs 300, 4 x 7 = 28 vs 20. KR = all minus line 8.
ATTIS_ALL, ATTIS_KR = (3000.0, 1200.0), (2000.0, 800.0)
SCHEME_ALL, SCHEME_KR = (3528.0, 1520.0), (2528.0, 1120.0)
HEALTHIUM_KR = (28.0, 20.0)                                                   # line 5
# Cash & TPA by hand. KR 101001 (bifurcated): molecule qty, all payers = 2 + 1 + 5 - 1 + 3
# (101001) + 1 (101003) = 11; allowed 0.1 x 11 = 1.1; cash/TPA qty 5 + 3 = 8; counted 6.9,
# shared 5 : 3 -> 4.3125 (IP, cash) and 2.5875 (OP, TPA). KR 201002: 6. KR 300001: 10.
# KR actual = 6.9 x 200 + 6 x 50 + 10 x 13 = 1810; potential = 6.9 x 600 + 6 x 120 + 10 x 28
# = 5140; loss 3330. DR 101001: molecule 1 + 2 = 3, allowed 0.3, cash 2, counted 1.7.
CASH_KR = {"actual_margin": 1810.0, "potential_margin": 5140.0,
           "used_mrp_value": 6.9 * 1200 + 600 + 200, "suggested_mrp_value": 6.9 * 1500 + 900 + 300}
CASH_ALL = (1810.0 + 1.7 * 200, 5140.0 + 1.7 * 600)


def write_sales(path: Path, lines, period="2026-08"):
    y, m = map(int, period.split("-"))
    wb = Workbook()
    ws = wb.active
    ws.title = "IP&OP Con Raw"
    ws.append(LINE_COLS)
    for i, ln in enumerate(lines):
        code = str(ln["code"])
        row = {"Type": ln["typ"], "LOCATIONNAME": ln["hosp"][1],
               "SALESINVOICENO": f"{ln['hosp'][0]}-{ln['typ']}-{700000 + i}",
               "SALESDATE": datetime(y, m, 6), "UHID": f"UH{i:06d}", "IP_OPNUMBER": f"IPN{i}",
               "BILLCATEGORY": "CREDIT" if ln["payer"] else "CASH", "PATNAME": f"Patient {i}",
               "TREATINGDOCID": ln["doc"][0], "TREATINGDOCNAME": ln["doc"][1],
               "MATERIALCODE": int(code), "MATERIALLINEITEM": f"ITEM {code}",
               "BRANDNAME": f"BRAND {code}", "QTY": ln["qty"], "MRP": 1.0,
               "TOTALMRP": ln["cost"] * 1.5, "COSTPRICE": 1.0, "TOTALCOSTPRICE": ln["cost"],
               "PAYER1": ln["payer"], "BATCHCODE": f"B{i}", "EXPIRYDATE": datetime(2028, 1, 1),
               "LOCATIONID": None, "GENERICNAME": f"GEN {code}", "ISSUECODE": f"PMIS{i}",
               "MANUFACTURERNAME": ln["maker"]}
        ws.append([row[c] for c in LINE_COLS])
    wb.save(path)


def write_scheme_list(path: Path, rows=SCHEME, fy="2026-27", sheet="SCHEME", cash=CASH):
    wb = Workbook()
    wb.active.title = "Sheet1"
    cs = wb.create_sheet("CASH&TPA")
    cs.append([f"FORMULARY FOR THE YEAR {fy} ( ONCO - 26-27 ,NON-ONCO 26-27 AND CONSUMBALES 26-27)"]
              + [None] * 13 + ["CASH PATIENTS --Suggestion Brand"])
    cs.append(CASH_HEADER)
    for used, up, mrp, gst, cat, share, sug, sp, smrp, sgst in cash:
        row = [None] * len(CASH_HEADER)
        row[0:9] = ["Onco", "MAKER", used, f"CASH NAME {used}", up, mrp, gst, cat, share]
        if sug:
            row[14:19] = [sug, f"CASH NAME {sug}", sp, smrp, sgst]
        cs.append(row)
    ws = wb.create_sheet(sheet)
    ws.append([f"FORMULARY FOR THE YEAR {fy} ( ONCO - 26-27 ,NON-ONCO 26-27 AND CONSUMBALES "
               f"26-27)"] + [None] * 14 + ["SCHEME PATIENTS --Suggestion Brand"])
    ws.append(SCHEME_HEADER)
    for used, up, cat, sug, sp, attis in rows:
        row = [None] * len(SCHEME_HEADER)
        row[0], row[2], row[3], row[4] = "Onco", "MAKER", used, f"LIST NAME {used}"
        row[1] = SCHEME_GROUP.get(used)
        row[5], row[6], row[9] = 1, up, cat
        if sug:
            row[15], row[16] = sug, f"LIST NAME {sug}"
            row[17] = "Attis Brand" if attis else "Other Brands"
            row[18] = sp
            row[19:23] = [0, 0, 0, 0]          # HCG's hand-filled month columns: never read
        ws.append(row)
    wb.save(path)


@pytest.fixture()
def env(tmp_path):
    dash = tmp_path / "raw" / "1829e863" / DASH
    (dash / "IP and OP sales").mkdir(parents=True)
    (dash / "Item Master Sheet").mkdir()
    write_sales(dash / "IP and OP sales" / "IP & OP Consolidated File August 26.xlsx", LINES)
    write_master(dash / "Item Master Sheet" / "Item master for formulary complience August -26.xlsx",
                 ITEM_MASTER)
    write_payer_master(dash / "Copy of Copy of Payer Master Final Sheet.xlsx")
    write_scheme_list(dash / "Copy of Onco Formulary -Spend Analysis Main Sheet_11.08.26.xlsx")
    ref = tmp_path / "reference"
    ref.mkdir()
    (ref / FC.HOSPITAL_MAP_FILE).write_text(HOSPITAL_MAP)
    (ref / FC.PAYER_OVERRIDES_FILE).write_text(OVERRIDES)
    (ref / C.BIFURCATION_FILE).write_text(BIFURCATION)
    dirs = dict(raw=tmp_path / "raw", cache=tmp_path / "cache", reference=ref)
    build_formulary(log=lambda *_: None, data=tmp_path / "formulary", **dirs)
    return dict(dirs=dirs, dash=dash, data=tmp_path / "spend", tmp=tmp_path)


def run(env, **kw):
    return build(log=lambda *_: None, data=env["data"], **env["dirs"], **kw)


def tables(env) -> dict:
    return {t: pd.read_parquet(env["data"] / f"{t}.parquet") for t in C.TABLES}


def used_sug(df) -> tuple:
    return (round(float(df["used_value"].sum()), 2), round(float(df["suggested_value"].sum()), 2))


# --- the scheme list -----------------------------------------------------------------------
def test_scheme_list_reads_hcg_layout(env):
    sl = load_scheme_list(next(env["dash"].glob("*Spend Analysis*")))
    assert sl.fy == (2026, 2027)
    assert sl.rows == len(SCHEME)
    assert set(sl.switches) == {"101001", "201002", "300001"}       # duplicate row merged
    a = sl.switches["101001"]
    assert (a.suggested, a.used_price, a.suggested_price, a.attis) == ("101901", 1000.0, 400.0, True)
    assert a.material_desc == "LIST NAME 101001" and a.suggested_desc == "LIST NAME 101901"
    assert sl.switches["201002"].hcg_category == "P-OPD"             # HCG's typo canonicalised
    assert sl.switches["201002"].attis is False
    assert [s["material"] for s in sl.skipped] == ["201004"]
    assert sl.covers("2026-04") and sl.covers("2027-03")
    assert not sl.covers("2026-03") and not sl.covers("2027-04")
    assert sl.switches["300001"].scheme_group == "SUTURES"


def test_cash_list_reads_hcg_layout(env):
    cl = load_cash_list(next(env["dash"].glob("*Spend Analysis*")))
    assert cl.fy == (2026, 2027) and cl.sheet == "CASH&TPA"
    assert cl.rows == len(CASH)
    assert set(cl.switches) == {"101001", "201002", "300001"}
    a = cl.switches["101001"]
    assert (a.suggested, a.used_price, a.used_mrp, a.used_gst, a.sharing) == ("101950", 1000.0, 1260.0, 5.0, 0.1)
    assert (a.used_margin, a.suggested_margin) == pytest.approx((200.0, 600.0))
    assert a.material_desc == "CASH NAME 101001" and a.suggested_desc == "CASH NAME 101950"
    assert [(s["material"], s["reason"]) for s in cl.skipped] == [("201004", "no used_gst")]


def test_bifurcation_file(env, tmp_path):
    b = load_bifurcation(env["dirs"]["reference"])
    assert b == {"101001": ("N", frozenset({"101001", "101003", "101950"}))}
    assert load_bifurcation(tmp_path / "no-such-folder") == {}
    bad = tmp_path / "bad"
    bad.mkdir()
    (bad / C.BIFURCATION_FILE).write_text("material,flag,group\n101001,N,101003 101950\n")
    with pytest.raises(BuildError, match="must include 101001 itself"):
        load_bifurcation(bad)
    (bad / C.BIFURCATION_FILE).write_text("material,flag,group\n101001,X,101001\n")
    with pytest.raises(BuildError, match="flag Y or N"):
        load_bifurcation(bad)


def test_conflicting_duplicate_aborts(tmp_path):
    p = tmp_path / "x Spend Analysis.xlsx"
    write_scheme_list(p, SCHEME + [(101001, 1000.0, "R", 101902, 380.0, True)])
    with pytest.raises(BuildError, match="listed twice"):
        load_scheme_list(p)


def test_missing_sheet_or_year_aborts(tmp_path):
    p = tmp_path / "a Spend Analysis.xlsx"
    write_scheme_list(p, sheet="SCHEME OLD")
    with pytest.raises(BuildError, match="no sheet 'SCHEME'"):
        load_scheme_list(p)
    p2 = tmp_path / "b Spend Analysis.xlsx"
    write_scheme_list(p2, fy="")
    with pytest.raises(BuildError, match="FORMULARY FOR THE YEAR"):
        load_scheme_list(p2)


# --- the numbers ---------------------------------------------------------------------------
def test_scheme_and_attis_values_by_hand(env):
    rep = run(env)
    sw = tables(env)[C.SWITCH]
    assert used_sug(sw[sw["attis"]]) == ATTIS_ALL
    assert used_sug(sw) == SCHEME_ALL
    kr = sw[sw["plant"] == "HC05"]
    assert used_sug(kr[kr["attis"]]) == ATTIS_KR
    assert used_sug(kr) == SCHEME_KR
    m = rep["months"]["2026-08"]
    assert m["attis"]["saving_missed"] == 1800.0 and m["attis"]["saving_pct"] == 60.0
    assert (m["kr"]["attis"]["used_value"], m["kr"]["attis"]["suggested_value"]) == ATTIS_KR
    assert (m["kr"]["scheme"]["used_value"], m["kr"]["scheme"]["suggested_value"]) == SCHEME_KR
    # only GOV lines with a priced suggestion: lines 1, 2, 4, 5, 8, 11
    assert int(sw["lines"].sum()) == 6
    assert set(sw["material"]) == {"101001", "201002", "300001"}
    assert m["gov"]["lines"] == 9                          # lines 1, 2, 4-9 and 11


def test_switch_rows_carry_prices_names_and_the_bill_cost(env):
    run(env)
    sw = tables(env)[C.SWITCH]
    d1 = sw[(sw["plant"] == "HC05") & (sw["doctor_id"] == "D1") & (sw["material"] == "101001")]
    assert len(d1) == 1                                      # line 1 and its return, line 11
    r = d1.iloc[0]
    assert (r["qty"], r["lines"], r["actual_cost"]) == (1.0, 2, 900.0)
    assert (r["used_price"], r["suggested_price"]) == (1000.0, 400.0)
    assert (r["suggested_material"], r["hcg_category"], bool(r["attis"])) == ("101901", "R", True)
    assert r["doctor_name"] == "Dr A" and r["category"] == "Onco Drugs"
    assert r["material_desc"] == "LIST NAME 101001"
    op = sw[(sw["patient_type"] == "OP")]
    assert list(op["doctor_id"]) == ["D2"]                  # GOV by payer_overrides.csv
    cons = sw[sw["material"] == "300001"].iloc[0]
    assert cons["category"] == "Consumables" and not cons["attis"]


def test_attis_consumption_every_payer(env):
    run(env)
    a = tables(env)[C.ATTIS_USE]
    assert set(a["material"]) == {"101901"}
    assert a["cost"].sum() == 1050.0 and int(a["lines"].sum()) == 2
    assert dict(zip(a["payer_bucket"], a["cost"])) == {"GOV": 700.0, "CASH": 350.0}
    assert dict(zip(a["plant"], a["patient_type"])) == {"HC05": "IP", "HC06": "OP"}


def test_nothing_patient_level_is_published(env):
    run(env)
    forbidden = FC.FORBIDDEN_COLUMNS | {"PAYER1", "MANUFACTURER", "SALES_DATE"}
    for name, df in tables(env).items():
        assert not {c.upper() for c in df.columns} & forbidden, name
        text = df.astype(str).to_csv()
        assert "UH0000" not in text and "Patient " not in text and "-700000" not in text, name


# --- cash & TPA -----------------------------------------------------------------------------
def cash_sums(df) -> tuple:
    return (round(float(df["actual_margin"].sum()), 2), round(float(df["potential_margin"].sum()), 2))


def test_cash_margin_loss_by_hand(env):
    rep = run(env)
    t = tables(env)
    cs = t[C.CASH_SWITCH]
    kr = cs[cs["plant"] == "HC05"]
    assert cash_sums(kr) == (CASH_KR["actual_margin"], CASH_KR["potential_margin"])
    assert cash_sums(cs) == pytest.approx(CASH_ALL)
    k = rep["months"]["2026-08"]["kr"]["cash"]
    assert k["margin_loss"] == 3330.0
    assert k["loss_pct"] == pytest.approx(3330 / 5140 * 100, abs=1e-4)          # loss / potential
    assert k["used_mrp_value"] == pytest.approx(CASH_KR["used_mrp_value"])
    assert k["topline_increase"] == pytest.approx(CASH_KR["suggested_mrp_value"] - CASH_KR["used_mrp_value"])
    assert set(cs["payer_bucket"]) == {"CASH", "CASH(TPA)"}                     # never GOV
    assert "101003" not in set(cs["material"])                                 # no suggestion


def test_bifurcation_counts_only_the_excess_and_shares_it_by_qty(env):
    run(env)
    t = tables(env)
    cs, al = t[C.CASH_SWITCH], t[C.CASH_ALLOWANCE]
    a = al[al["plant"] == "HC05"].iloc[0]
    assert (a["material"], a["flag"], a["sharing"], a["group_qty"], a["cash_qty"]) == ("101001", "N", 0.1, 11.0, 8.0)
    assert (a["allowed_qty"], a["excess_qty"]) == pytest.approx((1.1, 6.9))
    rows = cs[(cs["plant"] == "HC05") & (cs["material"] == "101001")].sort_values("patient_type")
    assert list(rows["cash_qty"]) == [5.0, 3.0]
    assert list(rows["qty"]) == pytest.approx([4.3125, 2.5875])
    assert list(rows["payer_bucket"]) == ["CASH", "CASH(TPA)"] and all(rows["bifurcated"])
    other = cs[cs["material"] == "201002"].iloc[0]
    assert (other["cash_qty"], other["qty"], bool(other["bifurcated"])) == (6.0, 6.0, False)
    d = al[al["plant"] == "HC06"].iloc[0]
    assert (d["group_qty"], d["excess_qty"]) == pytest.approx((3.0, 1.7))


def test_an_allowance_larger_than_the_use_counts_zero_never_negative(env):
    big = [(*r[:5], 0.9 if r[0] == 101001 else r[5], *r[6:]) for r in CASH]
    write_scheme_list(next(env["dash"].glob("*Spend Analysis*")), cash=big)
    run(env)
    cs = tables(env)[C.CASH_SWITCH]
    b = cs[cs["material"] == "101001"]
    assert float(b["qty"].sum()) == 0.0 and float(b["margin_loss"].sum()) == 0.0   # 9.9 > 8, 2.7 > 2


def test_without_the_bifurcation_file_every_cash_unit_counts(env):
    (env["dirs"]["reference"] / C.BIFURCATION_FILE).unlink()
    rep = run(env)
    k = rep["months"]["2026-08"]["kr"]["cash"]
    assert k["margin_loss"] == 8 * 400 + 6 * 70 + 10 * 15                        # 3770
    assert tables(env)[C.CASH_ALLOWANCE].empty
    assert rep["cash_list"]["bifurcation"]["brands"] == 0


def test_excluded_brands_are_not_counted_but_stay_in_molecule_groups(env):
    (env["dirs"]["reference"] / C.EXCLUDED_FILE).write_text(
        "material,material_desc,note\n201002,PARA,test\n101003,ONCO C,in a molecule group\n999999,GONE,not in the list\n")
    rep = run(env)
    cs = tables(env)[C.CASH_SWITCH]
    assert "201002" not in set(cs["material"])
    k = rep["months"]["2026-08"]["kr"]["cash"]
    assert k["margin_loss"] == 3330.0 - 6 * 70                    # 201002's 420 is gone
    al = tables(env)[C.CASH_ALLOWANCE]                            # 101003 still in 101001's group
    assert al.loc[al["plant"] == "HC05", "group_qty"].iloc[0] == 11.0
    ex = rep["cash_list"]["excluded"]
    assert ex["brands"] == ["201002"] and ex["not_in_cash_list"] == ["101003", "999999"]


def test_healthium_is_the_sutures_and_endo_scheme_rows(env):
    rep = run(env)
    h = rep["months"]["2026-08"]["kr"]["healthium"]
    assert (h["used_value"], h["suggested_value"]) == HEALTHIUM_KR
    sw = tables(env)[C.SWITCH]
    assert set(sw.loc[sw["scheme_group"] == "SUTURES", "material"]) == {"300001"}


# --- validation and publishing ---------------------------------------------------------------
def test_the_build_publishes_without_a_fixed_reference(env):
    # No typed figures gate the build: whatever the masters and sales give is published, and
    # KR's figures are only reported (tests/test_spend_golden.py sets them beside HCG's email).
    rep = run(env)
    assert "golden" not in rep["months"]["2026-08"]
    assert rep["months"]["2026-08"]["kr"]["attis"]["used_value"] == ATTIS_KR[0]
    assert (env["data"] / C.VALIDATION_FILE).exists()


def test_list_outside_the_financial_year_aborts(env):
    write_scheme_list(next(env["dash"].glob("*Spend Analysis*")), fy="2025-26")
    with pytest.raises(BuildError, match="FY 2025-26; it does not cover 2026-08"):
        run(env)


def test_empty_cache_aborts(env, tmp_path):
    with pytest.raises(BuildError, match="run python -m app.etl.formulary first"):
        build(log=lambda *_: None, raw=env["dirs"]["raw"], cache=tmp_path / "nothing",
              reference=env["dirs"]["reference"], data=env["data"])


def test_dry_run_publishes_nothing(env):
    run(env, dry_run=True)
    assert not env["data"].exists()
    assert (env["dirs"]["cache"] / "spend_validation.dry-run.json").exists()


def test_a_month_no_longer_cached_keeps_its_published_rows(env):
    run(env)
    sw = tables(env)[C.SWITCH]
    jul = sw.assign(period="2026-07")
    jul.to_parquet(env["data"] / f"{C.SWITCH}.parquet", index=False)      # pretend Jul was built
    run(env)
    after = tables(env)[C.SWITCH]
    assert sorted(after["period"].unique()) == ["2026-07", "2026-08"]
    assert used_sug(after[after["period"] == "2026-07"]) == SCHEME_ALL
    v = json.loads((env["data"] / C.VALIDATION_FILE).read_text())
    assert v["scheme_list"]["switches"] == 3 and v["scheme_list"]["attis_switches"] == 1
    assert v["scheme_list"]["skipped"][0]["material"] == "201004"


def test_cli_list_and_abort(env, monkeypatch, capsys):
    monkeypatch.setenv("FORMULARY_RAW_DIR", str(env["dirs"]["raw"]))
    monkeypatch.setenv("FORMULARY_CACHE_DIR", str(env["dirs"]["cache"]))
    monkeypatch.setenv("REFERENCE_DIR", str(env["dirs"]["reference"]))
    monkeypatch.setenv("SPEND_DATA_DIR", str(env["data"]))
    assert main(["--list"]) == 0
    out = capsys.readouterr().out
    assert "FY 2026-27" in out and "3 suggestions" in out and "2026-08" in out
    assert "cash list:" in out and "1 brands in cash_tpa_bifurcation.csv" in out
    assert "excluded:     0 cash brands left out per cash_tpa_excluded.csv" in out
    write_scheme_list(next(env["dash"].glob("*Spend Analysis*")), fy="2025-26")
    assert main([]) == 1
    assert "ABORT" in capsys.readouterr().err


# --- Task 4: blocks of interchangeable brands (P1/P2) ------------------------------------------
def test_brand_blocks_split_rank_and_flag_ties(tmp_path):
    from app.etl.spend.masters import load_brand_blocks
    p = tmp_path / "x Spend Analysis.xlsx"
    wb = Workbook()
    wb.active.title = "Sheet1"
    ws = wb.create_sheet("CASH&TPA")
    ws.append(["FORMULARY FOR THE YEAR 2026-27 ( ONCO )"])
    ws.append(CASH_HEADER)

    def brand(code, cat, share):
        row = [None] * len(CASH_HEADER)
        row[0:9] = ["Onco", "MAKER", code, f"NAME {code}", 10.0, 21.0, 5, cat, share]
        ws.append(row)

    brand(100001, "P", 0.3); brand(100002, "P-SCH", 0.7)
    ws.append(["Onco"])                                          # separator: no SAP CODE
    brand(100003, "R", 1)
    ws.append(["Onco", "ok"])                                    # HCG's 'ok' separator
    brand(100004, "P", 0.5); brand(100005, "P-SCH", 0.5); brand(100006, "OPD-R", 0); brand(100001, "P", 0.3)
    wb.save(p)
    b = load_brand_blocks(p).set_index("material")
    assert list(b.index) == ["100001", "100002", "100003", "100004", "100005", "100006"]   # dup kept once
    assert list(b["block_id"]) == [1, 1, 2, 3, 3, 3]
    assert list(b["p_rank"]) == [2, 1, 1, 1, 1, 3]
    assert list(b["tie"]) == [False, False, False, True, True, False]
    assert list(b["block_size"]) == [2, 2, 1, 3, 3, 3] and set(b["block_sharing_sum"]) == {1.0}
    assert b.loc["100006", "ms_category"] == "R-OPD"                                        # typo fixed


def test_the_build_publishes_brand_blocks_whole(env):
    rep = run(env)
    bb = tables(env)[C.BRAND_BLOCKS]
    assert "period" not in bb.columns and len(bb) == 5                     # the 5 coded CASH rows
    a = bb.set_index("material").loc["101001"]
    assert (a["cash_suggested"], a["scheme_suggested"], bool(a["scheme_attis"])) == ("101950", "101901", True)
    assert rep["brand_blocks"]["brands"] == 5
    run(env)                                                               # rebuild: replaced, not appended
    assert len(tables(env)[C.BRAND_BLOCKS]) == 5
