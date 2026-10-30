# tests/test_purchase_readers.py — the purchase build's readers on tiny made-up workbooks.
#
# HCG's GRN extracts and out-of-formulary workbooks never enter the repo, so each test writes
# a miniature copy of their layout in tmp_path with openpyxl: padded and misspelt headers, an
# unnamed column, GR dates as real dates, Excel serials and 'dd.mm.yyyy' text, lines without a
# material, returns, GR / PO / invoice numbers that must never be read, HCG's worksheet,
# parked and Summary sheets. The .xlsb layouts (pyxlsb: codes and dates as floats, '0x2a'
# error cells) go through the same row parsers as plain Python lists. Every directory is
# passed explicitly: nothing here touches app/data or the real cache.
from __future__ import annotations

import os
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from app.etl.formulary.discover import BuildError
from app.etl.purchase import config as C
from app.etl.purchase import grn as G
from app.etl.purchase import masters as M
from tests.test_formulary_pipeline import DASH, write_master
from tests.test_spend_pipeline import CASH, write_scheme_list

HASH = "1829e863"
GRN_DIR = "GRN Dec 25 to May 2026"
OOF_DIR = "Out of Formulary Spend Analysis @6-27"
OOF_MASTER_DIR = "Out Of formulary Master Sheet (Out of formulary Brands vs Proposed Brands to use)"
OOF_MASTER_FILE = "Copy of Out of formualry vs suggested formualry brand master-2026-27.xlsx"

# Identifiers sitting in the source columns: none may reach a frame, a cache file or a report.
SECRETS = ("4500999001", "5000777001", "INV-SECRET-9", "REF-SECRET-9")

# HCG's GRN header in HCG's order and spelling (padded names, 'mvt type', an unnamed column).
GRN_HEADER = [
    "Sl.No", "Company Code", "PR No", "PO No", "PO Date", "GR No", "GR Li item", "mvt type",
    "reference document", "GR Date", "PO to GR TAT", "PR to GR TAT", "Goods supplier",
    "Goods supplier name", "Vendor Inv No", "Plant", "S.Loc", "Material", "Material Description",
    "Batch", None, "Expiry Date", "GR Qty", "GR UOM", "    Base Price", " Dicsount",
    "Discount Value", "     Net Price", "Total Amount without tax", " Tax %", "   Tax Amount",
    "Total Amount with Tax", "      Unit MRP", "Total MRP Value w/o free", "Manufacturer",
    "Material Type", "Minor Category", "Major Group", "Minor Group", "Formulary"]


def serial(d: date) -> float:
    return float((d - date(1899, 12, 30)).days)


def records(df: pd.DataFrame) -> list:
    """Rows as lists, missing values as None."""
    return df.astype(object).where(df.notna(), None).values.tolist()


def R(mvt, gr_date, plant, sloc, material, total, qty=1.0, base=None, mrp=0.0):
    """One GRN line: codes written as floats like SAP's export, an identifier in every
    identifier column."""
    base = total / qty if base is None else base
    return {"Sl.No": 1, "Company Code": 1010, "PR No": None, "PO No": float(SECRETS[0]),
            "PO Date": "01.06.2026", "GR No": float(SECRETS[1]), "GR Li item": 1, "mvt type": mvt,
            "reference document": SECRETS[3] if mvt != 101 else None, "GR Date": gr_date,
            "PO to GR TAT": 3, "PR to GR TAT": None, "Goods supplier": 300100.0,
            "Goods supplier name": "MEDS DISTRIBUTOR", "Vendor Inv No": SECRETS[2],
            "Plant": plant, "S.Loc": sloc, "Material": float(material) if material else None,
            "Material Description": f"ITEM {material}", "Batch": "B1", None: None,
            "Expiry Date": datetime(2028, 1, 31), "GR Qty": qty, "GR UOM": "EA",
            "    Base Price": base, " Dicsount": 0, "Discount Value": 0.0,
            "     Net Price": base, "Total Amount without tax": total, " Tax %": 5.0,
            "   Tax Amount": total * 0.05, "Total Amount with Tax": total * 1.05,
            "      Unit MRP": mrp, "Total MRP Value w/o free": mrp * qty, "Manufacturer": "ACME",
            "Material Type": "ZNOC-Medical Non Onco Drugs", "Minor Category": "DRUG",
            "Major Group": "DRUGS", "Minor Group": "MOLECULE", "Formulary": "FORMULARY"}


BLANK_ROW = {"Material Description": "   "}                  # only spaces: an empty row
# June by hand (Total Amount without tax in the 6th field). Funnel in HCG's order:
#   read       all 9 lines (8 + 1 without material)     value 1414
#   mvt 101    drops the two returns                  7 lines, 1429
#   series     drops the consumable and the no-material line   5, 370
#   S125       drops line 2                           4, 340
#   BH*        drops line 4                           3, 300
#   HM03/HM04  drops line 5                           2, 250
JUNE = [
    R(101, "04.06.2026", "HC05", "S124", 100001, 200.0, qty=2.0, mrp=150.0),   # 1 kept
    R(101, datetime(2026, 6, 5), "HC05", "S125", 200002, 30.0),               # 2 S125
    R(122, "05.06.2026", "HC05", "S124", 200002, -10.0),                      # 3 return
    R(101, serial(date(2026, 6, 10)), "BH03", "S124", 200003, 40.0),          # 4 Milann
    R(101, "10.06.2026", "HM03", "S124", 200004, 50.0),                       # 5 HM03
    R(101, "11.06.2026", "HC06", "S124", 300001, 60.0),                       # 6 consumable
    R(101, "12.06.2026", "HC06", "S124", None, 999.0),                        # 7 no material
    BLANK_ROW,
    R(161, "13.06.2026", "HC06", "S124", 100001, -5.0),                       # 9 return
    R(101, "01.06.2026", "HC06", "S124", 200005, 50.0),                       # 10 kept
]
JULY = [R(101, "02.07.2026", "HC05", "S124", 200006, 70.0),
        R(101, datetime(2026, 7, 9), "HC06", "S124", 100001, 80.0)]


def write_grn(path: Path, lines, sheet="GRN report", scratch=True):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(GRN_HEADER)
    for ln in lines:
        ws.append([ln.get(h) for h in GRN_HEADER])
    if scratch:                                          # an analyst's scratch sheet: ignored
        wb.create_sheet("Sheet1").append(["Company Code", 1010])
    wb.save(path)


@pytest.fixture()
def grn_env(tmp_path):
    folder = tmp_path / "raw" / HASH / GRN_DIR
    folder.mkdir(parents=True)
    write_grn(folder / "03.Consolidated GRN June 2026.xlsx", JUNE)
    write_grn(folder / "04.Consolidated GRN July 2026.xlsx", JULY)
    return dict(raw=tmp_path / "raw", cache=tmp_path / "cache", folder=folder, tmp=tmp_path)


def read(env, **kw):
    return G.read_grn_files(raw=env["raw"], cache=env["cache"], log=lambda *_: None, **kw)


# --- dates and codes -----------------------------------------------------------------------
@pytest.mark.parametrize("value, expected", [
    (datetime(2026, 8, 4, 10, 30), date(2026, 8, 4)),
    (date(2026, 8, 4), date(2026, 8, 4)),
    (pd.Timestamp("2026-08-04"), date(2026, 8, 4)),
    (serial(date(2026, 8, 4)), date(2026, 8, 4)),          # pyxlsb: a float
    (int(serial(date(2026, 8, 4))), date(2026, 8, 4)),
    ("04.08.2026", date(2026, 8, 4)),                       # day first: 4 August, not 8 April
    ("4.8.2026", date(2026, 8, 4)),
    ("04-08-2026", date(2026, 8, 4)),
    ("04/08/2026 00:00", date(2026, 8, 4)),
    (" 31.08.2026 ", date(2026, 8, 31)),
    ("2026-08-04", date(2026, 8, 4)),
    ("2026-08-04 00:00:00", date(2026, 8, 4)),
    ("08.13.2026", None),                                   # no 13th month: never swapped
    ("31.02.2026", None),
    ("04.08.26", None),                                     # two-digit year: not guessed
    ("04.08-2026", None),
    ("46238", None),                                        # a number as text
    ("August 4", None),
    (5.0, None),                                            # a serial outside 1990-2100
    (True, None),
    (None, None), ("", None), (float("nan"), None), (pd.NaT, None),
])
def test_parse_gr_date_reads_day_first_and_never_guesses(value, expected):
    assert G.parse_gr_date(value) == expected


def test_parse_gr_dates_counts_blank_and_unreadable_cells():
    p = G.parse_gr_dates(["04.08.2026", "04.08.2026", datetime(2026, 8, 5), serial(date(2026, 8, 6)),
                          None, "  ", "31.02.2026", "n/a"])
    assert p.dates[:4] == [date(2026, 8, 4), date(2026, 8, 4), date(2026, 8, 5), date(2026, 8, 6)]
    assert p.dates[4:] == [None] * 4
    assert (p.blank, p.unparseable) == (2, 2)
    assert p.kinds == {"date": 1, "serial": 1, "text": 4}


def test_single_period_needs_exactly_one_month():
    assert G.single_period([date(2026, 8, 1), None, date(2026, 8, 31)], "x") == "2026-08"
    with pytest.raises(BuildError, match="span 2 months"):
        G.single_period([date(2026, 8, 1), date(2026, 9, 1)], "x")
    with pytest.raises(BuildError, match="no readable GR date"):
        G.single_period([None, None], "x")


def test_clean_code_never_makes_nan_text():
    assert [G.clean_code(v) for v in (100301.0, "100301.0", " 100301 ", 101, "S125")] == \
        ["100301", "100301", "100301", "101", "S125"]
    assert [G.clean_code(v) for v in (None, "", "  ", float("nan"), pd.NA, "#N/A", "0x2a")] == [None] * 7
    s = G.clean_codes(pd.Series([100301.0, None, float("nan"), "nan-free 7"], dtype="object"))
    assert s.isna().tolist() == [False, True, True, False]
    assert s[0] == "100301" and "nan" not in s.dropna().str.lower().tolist()


# --- one GRN workbook ----------------------------------------------------------------------
def test_read_grn_file_normalises_headers_dates_and_drops_blank_materials(grn_env):
    df, rep = G.read_grn_file(grn_env["folder"] / "03.Consolidated GRN June 2026.xlsx",
                              cache=grn_env["cache"], log=lambda *_: None)
    assert list(df.columns) == ["period", *C.GRN_COLUMNS]
    assert (rep.file, rep.sheet, rep.period) == ("03.Consolidated GRN June 2026.xlsx", "GRN report", "2026-06")
    assert (rep.rows, rep.lines, rep.blank_rows) == (9, 8, 1)
    assert rep.blank_material == {"101": {"lines": 1, "value": 999.0}}
    assert rep.blank_material_lines == 1
    assert rep.mvt == {"101": 7, "122": 1, "161": 1} and rep.returns == 2
    assert rep.dates == {"date": 1, "serial": 1, "text": 7, "blank": 0, "unparseable": 0}
    assert rep.value == pytest.approx(200 + 30 - 10 + 40 + 50 + 60 - 5 + 50)
    assert not rep.cached
    assert df["material"].tolist() == ["100001", "200002", "200002", "200003", "200004", "300001",
                                       "100001", "200005"]
    assert df["mvt_type"].tolist() == ["101", "101", "122", "101", "101", "101", "161", "101"]
    assert df["vendor_code"].unique().tolist() == ["300100"]
    assert df["gr_date"].dt.day.tolist() == [4, 5, 5, 10, 10, 11, 13, 1]      # never month-first
    assert (df["period"] == "2026-06").all() and df["gr_date"].dt.month.eq(6).all()
    first = df.iloc[0]
    assert (first["plant"], first["sloc"], first["uom"], first["base_price"], first["unit_mrp"],
            first["gr_qty"], first["tax_pct"], first["grn_formulary"]) == \
        ("HC05", "S124", "EA", 100.0, 150.0, 2.0, 5.0, "FORMULARY")
    assert str(df["gr_qty"].dtype) == "float64" and str(df["gr_date"].dtype).startswith("datetime64")


def test_no_identifier_leaves_the_reader(grn_env):
    read_ = read(grn_env)
    frame_text = read_.frame.astype(str).to_csv()
    files = list(grn_env["cache"].iterdir())
    assert {p.suffix for p in files} == {".parquet", ".json"}
    blobs = [p.read_bytes() for p in files] + [frame_text.encode(), str(read_.report()).encode()]
    for secret in SECRETS:
        assert not any(secret.encode() in b for b in blobs), secret
    forbidden = {"gr_no", "po_no", "pr_no", "invoice", "reference", "batch", "vendor_inv"}
    cols = set(pd.read_parquet(next(grn_env["cache"].glob("*.parquet"))).columns)
    assert not {c for c in cols if any(f in c for f in forbidden)}


def test_the_cache_is_reused_and_refreshed_when_the_file_changes(grn_env):
    path = grn_env["folder"] / "04.Consolidated GRN July 2026.xlsx"
    df1, rep1 = G.read_grn_file(path, cache=grn_env["cache"], log=lambda *_: None)
    df2, rep2 = G.read_grn_file(path, cache=grn_env["cache"], log=lambda *_: None)
    assert (rep1.cached, rep2.cached) == (False, True)
    pd.testing.assert_frame_equal(df1, df2)
    write_grn(path, JULY + [R(101, "20.07.2026", "HC05", "S124", 200007, 5.0)])
    st = path.stat()
    os.utime(path, ns=(st.st_atime_ns, st.st_mtime_ns + 10**9))
    df3, rep3 = G.read_grn_file(path, cache=grn_env["cache"], log=lambda *_: None)
    assert not rep3.cached and len(df3) == 3
    assert len(list(grn_env["cache"].glob("grn_*.parquet"))) == 1     # the stale copy is gone


def test_extra_columns_are_read_on_request(grn_env):
    path = grn_env["folder"] / "04.Consolidated GRN July 2026.xlsx"
    df, _ = G.read_grn_file(path, cache=grn_env["cache"], extra=("po_date", "tax_amount"),
                            log=lambda *_: None)
    assert list(df.columns)[-2:] == ["po_date", "tax_amount"]
    assert df["po_date"].dt.date.tolist() == [date(2026, 6, 1)] * 2      # 'dd.mm.yyyy', day first
    assert df["tax_amount"].tolist() == pytest.approx([3.5, 4.0])
    with pytest.raises(ValueError, match="unknown extra GRN columns"):
        G.read_grn_file(path, cache=grn_env["cache"], extra=("po_no",))


# --- every GRN workbook --------------------------------------------------------------------
def test_read_grn_files_assigns_each_workbook_to_its_month(grn_env):
    r = read(grn_env)
    assert r.folder.name == GRN_DIR and r.periods == ["2026-06", "2026-07"]
    assert [f.period for f in r.files] == ["2026-06", "2026-07"]
    assert r.frame.groupby("period").size().to_dict() == {"2026-06": 8, "2026-07": 2}
    assert r.blank_lines().to_dict("records") == [
        {"period": "2026-06", "mvt_type": "101", "lines": 1, "value": 999.0}]
    assert [f["file"] for f in r.report()["files"]] == ["03.Consolidated GRN June 2026.xlsx",
                                                         "04.Consolidated GRN July 2026.xlsx"]
    assert all(f.cached for f in read(grn_env).files)
    only = read(grn_env, months=["2026-07"])
    assert only.periods == ["2026-07"] and len(only.frame) == 2
    with pytest.raises(BuildError, match="no GRN workbook for 2026-09"):
        read(grn_env, months=["2026-09"])


def test_unreadable_dates_are_counted_not_guessed(grn_env):
    write_grn(grn_env["folder"] / "04.Consolidated GRN July 2026.xlsx",
              JULY + [R(101, "31.02.2026", "HC05", "S124", 200008, 1.0)])
    r = read(grn_env)
    july = r.files[1]
    assert july.period == "2026-07" and july.unparseable_dates == 1 and july.lines == 3
    assert r.frame.loc[r.frame["material"] == "200008", "gr_date"].isna().all()
    assert r.warnings == ["04.Consolidated GRN July 2026.xlsx: 1 GR dates could not be read "
                          "(kept, month from the file)"]


def test_a_workbook_spanning_two_months_aborts(grn_env):
    write_grn(grn_env["folder"] / "04.Consolidated GRN July 2026.xlsx",
              JULY + [R(101, "01.08.2026", "HC05", "S124", 200008, 1.0)])
    with pytest.raises(BuildError, match=r"span 2 months \(2026-07: 2, 2026-08: 1\)"):
        read(grn_env)


def test_two_workbooks_for_one_month_abort(grn_env):
    write_grn(grn_env["folder"] / "05.GRN extra.xlsx", JULY)
    with pytest.raises(BuildError, match="two GRN workbooks hold 2026-07"):
        read(grn_env)


def test_a_workbook_without_the_grn_columns_aborts(grn_env):
    wb = Workbook()
    wb.active.append([h for h in GRN_HEADER if h != "GR Date"])
    wb.save(grn_env["folder"] / "05.GRN broken.xlsx")
    with pytest.raises(BuildError, match="no sheet has the GRN columns.*lacks GR DATE"):
        read(grn_env)


def test_the_grn_folder_must_be_unique(grn_env):
    (grn_env["raw"] / HASH / "GRN Sept 2026").mkdir()
    with pytest.raises(BuildError, match="more than one GRN folder"):
        read(grn_env)


# --- scope ---------------------------------------------------------------------------------
def test_apply_scope_follows_hcg_order_and_counts_lines_without_material(grn_env):
    r = read(grn_env)
    kept, funnel = G.apply_scope(r.frame, C.load_scope(), r.blank_lines())
    june = funnel[funnel["period"] == "2026-06"]
    assert june["kind"].tolist() == ["lines", "mvt_keep", "series_keep", "sloc", "plant_prefix", "plant"]
    assert june["label"].tolist() == ["lines in the GRN", "movement type 101", "material series 1, 2",
                                      "not storage location S125", "not plants BH*",
                                      "not plants HM03, HM04"]
    assert june["lines"].tolist() == [9, 7, 5, 4, 3, 2]
    assert june["value"].tolist() == pytest.approx([1414, 1429, 370, 340, 300, 250])
    assert funnel[funnel["period"] == "2026-07"]["lines"].tolist() == [2] * 6
    assert sorted(kept["material"]) == ["100001", "100001", "200005", "200006"]
    # without the reader's report, the first steps miss the line that has no material
    _, plain = G.apply_scope(r.frame, C.load_scope())
    assert plain[plain["period"] == "2026-06"]["lines"].tolist() == [8, 6, 5, 4, 3, 2]


def test_the_committed_scope_is_hcgs(tmp_path):
    s = C.load_scope()
    assert (s.mvt_keep, s.series_keep, s.sloc, s.plant_prefix, s.plant) == (
        {"101"}, {"1", "2"}, {"S125"}, ("BH",), {"HM03", "HM04"})
    assert all(r["reason"] for r in s.report())


@pytest.mark.parametrize("body, error", [
    ("kind,value,reason\nmvt_keep,101,x\nseries_keep,2,x\nplants,HM03,x\n", "kind must be one of"),
    ("kind,value,reason\nmvt_keep,101,x\nseries_keep,,x\n", "needs a value"),
    ("kind,value,reason\nmvt_keep,101,x\nseries_keep,21,x\n", "one digit"),
    ("kind,value,reason\nmvt_keep,101,x\nmvt_keep,101.0,x\nseries_keep,2,x\n", "listed twice"),
    ("kind,value,reason\nsloc,S125,x\nseries_keep,2,x\n", "kind mvt_keep"),
    (None, "scope file not found"),
])
def test_bad_scope_files_are_rejected(tmp_path, body, error):
    if body is not None:
        (tmp_path / C.SCOPE_FILE).write_text(body)
    with pytest.raises(BuildError, match=error):
        C.load_scope(tmp_path)


def test_a_scope_file_with_its_own_values(tmp_path):
    (tmp_path / C.SCOPE_FILE).write_text("kind,value,reason\nmvt_keep,101,x\nseries_keep,3,x\n"
                                         " PLANT_PREFIX , hm ,x\n")
    s = C.load_scope(tmp_path)
    assert (s.series_keep, s.plant_prefix, s.sloc, s.plant) == ({"3"}, ("HM",), frozenset(), frozenset())


# --- HCG's OOF master (xlsb: rows as pyxlsb gives them) -----------------------------------
MASTER_ROWS = [
    ["OUT OF FORMULARY BRANDS", None, None, None, "SUGESTED FORMULARY BRANDS ", None, None, None, None],
    ["Material", "Material Description", "Generic name", None, "MATERIAL", "DESCRIPTION", "PRICE",
     "MRP", "GST"],
    [205411.0, "LOBET 100MG TAB 10'S", "LABETALOL", None, 227432.0, "LAB 100MG", 12.907626666666667,
     174.38, 5.0],
    [217038.0, "LACTIHEP PLUS SYR", "LACTULOSE", None, 227060.0, "LAXSIA PLUS SYP", 77.0, 174.38, 5.0],
    [217055.0, "CREMOPEN PLUS SYR", "LACTULOSE", None, 227060.0, "LAXSIA PLUS SYP", 94.8152, 174.38, 5.0],
    [205411.0, "LOBET 100MG TAB 10'S", "LABETALOL", None, 227432.0, "LAB 100MG", 12.907626666666667,
     174.38, 5.0],                                               # repeated as is: kept once
    [200001.0, "NO GST BRAND", "X", None, 210001.0, "Y", 1.0, 2.0, "0x2a"],   # #N/A GST: skipped
    [None, None, None, None, None, None, None, None, None],
]


def test_oof_master_rows_parse_hcg_layout():
    m = M.parse_oof_master_rows(MASTER_ROWS, name=OOF_MASTER_FILE.replace(".xlsx", ".xlsb"))
    assert (m.rows, len(m.entries), m.fy) == (5, 3, (2026, 2027))
    assert m.skipped == [{"row": 7, "material": "200001", "material_desc": "NO GST BRAND",
                          "reason": "no gst"}]
    e = m.entries["205411"]
    assert (e.row, e.suggested, e.suggested_desc, e.generic_name, e.mrp, e.gst) == \
        (3, "227432", "LAB 100MG", "LABETALOL", 174.38, 5.0)
    # one suggested code, two prices: the price belongs to the OOF row
    assert (m.entries["217038"].price, m.entries["217055"].price) == (77.0, 94.8152)
    f = m.frame()
    assert list(f.columns) == ["row", "material", "material_desc", "generic_name", "suggested",
                               "suggested_desc", "price", "mrp", "gst"] and len(f) == 3
    assert m.covers("2026-08") and m.covers("2027-03") and not m.covers("2027-04")


def test_oof_master_conflicting_rows_abort():
    rows = MASTER_ROWS[:3] + [[205411.0, "LOBET", "LABETALOL", None, 227999.0, "OTHER", 1.0, 2.0, 5.0]]
    with pytest.raises(BuildError, match="205411 listed twice"):
        M.parse_oof_master_rows(rows)


# --- HCG's monthly OOF workbooks (xlsx written here; xlsb rows below) ----------------------
WS_HEADER = [
    "GR No", "GR Date", "Goods supplier", "Goods supplier name", "Plant", "S.Loc", "Material",
    "Material Description", "GR Qty", "    Base Price", "Total Amount without tax", " Tax %",
    "      Unit MRP", "Manufacturer", "Minor Group", "Formulary Flagging ", "Net price", "Abb mrp",
    "Unit margin", "TOTAL PP", "TOTAL MRP", "Unit margin %", "Plant", "SAP code", "SAP NAME",
    "HCG Pharma Selling Price(Rs) Excl of VAT", "MRP(Rs) Incl of GST", "GST %", "Abb MRP",
    "Unit Margin Value", "Unit Margin", "Total PP", "Total MRP"]


def W(gr_date, plant, material, qty, base, sug=None, price=None, mrp=None, gst=5.0):
    """One worksheet line in HCG's A..AG layout; sug None = HCG's #N/A lookup."""
    row = [float(SECRETS[1]), gr_date, 300100.0, "MEDS DISTRIBUTOR", plant, "S124", float(material),
           f"ITEM {material}", qty, base, qty * base, 5.0, base * 2, "ACME", "MOLECULE",
           "OUT OF FORMULARY"] + [None] * 6 + [plant]
    if sug:
        row += [float(sug), f"SUG {sug}", price, mrp, gst]
    else:
        row += ["#N/A"] * 5
    return row + [None] * 5


AUG_WS = [
    W("04.08.2026", "HC05", 200010, 10.0, 5.0, 210010, 3.0, 10.5),
    W(datetime(2026, 8, 5), "HC05", 200010, 4.0, 5.0, 210010, 3.0, 10.5),     # same pair: summed
    W("06.08.2026", "HC05", 200010, 1.0, 6.0, 210010, 3.0, 10.5),             # other base price
    W("07.08.2026", "hc06", 100020, 2.0, 100.0, 110020, 50.0, 210.0),
    W("08.08.2026", "HC06", 200030, 3.0, 7.0),                                # HCG could not price
]
SUMMARY_HEADER = ["Center", "Center Name ", "Out of formulary used brand PP",
                  "Out of formulary used brand MRP", "Margi Value", "Margin %", None,
                  "Suggesed formulary brand PP", "Suggesed formulary brand MRP", "Margi Value",
                  "Margin %", "Margin Value Differnce "]
# HC05 typed as values; HC06 as formulas openpyxl saves without a cached value.
SUMMARY_ROWS = [
    ["HC05", "KR", 100.0, 150.0, 50.0, 0.33, None, 80.0, 200.0, 120.0, 0.6, 70.0],
    ["HC06", "DR", 10.0, 12.0, "=D4-C4", "=E4/D4", None, 8.0, 20.0, "=I4-H4", "=J4/I4", "=J4-E4"],
    ["Grand Total", None, 110.0, 162.0, None, None, None, 88.0, 220.0, None, None, None],
]


def write_hcg_workbook(path: Path, ws_sheets: dict, month="July", summary_sheet="Summary",
                       summary_rows=SUMMARY_ROWS):
    """HCG's layout; the Summary titles name `month` (HCG's August file says July)."""
    wb = Workbook()
    wb.active.title = "Consolidated GRN August 26"          # HCG's copy of the GRN: never read
    wb.active.append(GRN_HEADER)
    for name, lines in ws_sheets.items():
        ws = wb.create_sheet(name)
        ws.append(["OUT OF FORMULARY PURCHASE August 2026"] + [None] * 21 + ["SUGGESTED FORMULARY BRANDS"])
        ws.append(WS_HEADER)
        for ln in lines:
            ws.append(ln)
    wb.create_sheet("Pivote").append(["Row Labels", "Sum of TOTAL PP"])
    s = wb.create_sheet(summary_sheet)
    s.append([f"ACTUAL OUT OF FORMULARY PURCHASE VALUE {month} 2026"] + [None] * 6
             + [f"SUGGESTED FORMULARY BRANDS VALUE {month} 26"])
    s.append(SUMMARY_HEADER)
    for r in summary_rows:
        s.append(r)
    wb.save(path)


@pytest.fixture()
def oof_env(tmp_path):
    d = tmp_path / "raw" / HASH / OOF_DIR
    (d / OOF_MASTER_DIR).mkdir(parents=True)
    wb = Workbook()
    for r in MASTER_ROWS:
        wb.active.append(r)
    wb.save(d / OOF_MASTER_DIR / OOF_MASTER_FILE)
    write_hcg_workbook(d / "Out of Formulary Brand Margin loss Spend anlayiss August 26.xlsx",
                       {"Out of formulary August 26 WS": AUG_WS})
    return dict(raw=tmp_path / "raw", dir=d)


def test_hcg_workbook_month_comes_from_gr_dates_not_titles(oof_env):
    wb = M.load_hcg_workbook(oof_env["dir"] / "Out of Formulary Brand Margin loss Spend anlayiss August 26.xlsx")
    assert (wb.period, wb.ws_sheet, wb.ws_lines, wb.ws_unpriced) == \
        ("2026-08", "Out of formulary August 26 WS", 5, 1)
    assert any("2026-07" in w and "2026-08 used" in w for w in wb.warnings)
    assert wb.parked_sheet is None and wb.parked.empty and list(wb.parked.columns) == M.PARKED_COLUMNS
    p = wb.pairs
    assert list(p.columns) == M.PAIR_COLUMNS
    assert records(p[["plant", "material", "base_price", "suggested", "lines", "qty"]]) == [
        ["HC05", "200010", 5.0, "210010", 2, 14.0],
        ["HC05", "200010", 6.0, "210010", 1, 1.0],
        ["HC06", "100020", 100.0, "110020", 1, 2.0],
        ["HC06", "200030", 7.0, None, 1, 3.0]]
    assert p.iloc[0][["suggested_desc", "sug_price", "sug_mrp", "sug_gst", "material_desc"]].tolist() == \
        ["SUG 210010", 3.0, 10.5, 5.0, "ITEM 200010"]
    assert pd.isna(p.iloc[3]["sug_price"])
    assert not any(SECRETS[1] in str(v) for v in p.astype(str).values.ravel())


def test_summary_loss_is_recomputed_when_a_formula_has_no_cached_value(oof_env):
    wb = M.load_hcg_workbook(next(oof_env["dir"].glob("*August 26.xlsx")))
    s = wb.summary
    assert list(s.columns) == M.SUMMARY_COLUMNS and (s["period"] == "2026-08").all()
    assert s.set_index("plant")[["centre_name", "used_margin", "sug_margin", "loss"]].to_dict("index") == {
        "HC05": {"centre_name": "KR", "used_margin": 50.0, "sug_margin": 120.0, "loss": 70.0},
        "HC06": {"centre_name": "DR", "used_margin": 2.0, "sug_margin": 12.0, "loss": 10.0}}
    assert wb.summary_computed == 1
    assert wb.report()["loss"] == 80.0 and wb.report()["centres"] == 2


def test_the_final_sheet_wins_and_parked_lines_are_read(tmp_path):
    june_ws = [W("02.06.2026", "HC05", 200010, 1.0, 5.0, 210010, 3.0, 10.5),
               W("03.06.2026", "HC05", 200040, 1.0, 9.0)]
    final = june_ws[:1]
    parked = [W("09.06.2026", "HC16", 205013, 20.0, 1.5), W("10.06.2026", "HC16", 205013, 10.0, 1.5)]
    path = tmp_path / "Out of Formulary Spend Analysis June 26 WS.xlsx"
    write_hcg_workbook(path, {"Out of Formulary WS June 26": june_ws, "WS Final Sheet": final,
                              "Need To check flaging later": parked},
                       month="June", summary_sheet="Summary ")
    wb = M.load_hcg_workbook(path)
    assert (wb.period, wb.ws_sheet, wb.ws_lines, wb.summary_sheet) == ("2026-06", "WS Final Sheet", 1, "Summary ")
    assert wb.parked_sheet == "Need To check flaging later"
    assert wb.parked.to_dict("records") == [{"period": "2026-06", "plant": "HC16", "material": "205013",
                                             "material_desc": "ITEM 205013", "base_price": 1.5,
                                             "lines": 2, "qty": 30.0}]
    assert not wb.warnings


def test_parked_lines_of_another_month_abort(tmp_path):
    path = tmp_path / "Out of Formulary Spend Analysis July 26 WS.xlsx"
    write_hcg_workbook(path, {"Out of formulary July 26 WS": [W("02.07.2026", "HC05", 200010, 1.0, 5.0,
                                                                210010, 3.0, 10.5)],
                              "Need To check flaging later": [W("09.06.2026", "HC16", 205013, 1.0, 1.5)]})
    with pytest.raises(BuildError, match="parked lines are 2026-06, the worksheet 2026-07"):
        M.load_hcg_workbook(path)


def test_hcg_files_are_found_and_keyed_by_month(oof_env):
    write_hcg_workbook(oof_env["dir"] / "Out of Formulary Spend Analysis July 26 WS.xlsx",
                       {"Out of formulary July 26 WS": [W("02.07.2026", "HC05", 200010, 1.0, 5.0,
                                                          210010, 3.0, 10.5)]})
    assert [p.name for p in M.find_hcg_workbooks(oof_env["raw"])] == [
        "Out of Formulary Brand Margin loss Spend anlayiss August 26.xlsx",
        "Out of Formulary Spend Analysis July 26 WS.xlsx"]
    assert list(M.load_hcg_workbooks(oof_env["raw"])) == ["2026-07", "2026-08"]
    assert M.find_oof_master(oof_env["raw"]).name == OOF_MASTER_FILE
    m = M.load_oof_master(raw=oof_env["raw"])
    assert (m.name, m.sheet, len(m.entries), m.fy) == (OOF_MASTER_FILE, "Sheet", 3, (2026, 2027))
    write_hcg_workbook(oof_env["dir"] / "Out of Formulary Spend Analysis July 26 copy.xlsx",
                       {"Out of formulary July 26 WS": [W("03.07.2026", "HC05", 200010, 1.0, 5.0)]})
    with pytest.raises(BuildError, match="two OOF workbooks hold 2026-07"):
        M.load_hcg_workbooks(oof_env["raw"])


def test_a_workbook_without_a_worksheet_or_summary_aborts():
    with pytest.raises(BuildError, match="needs one OOF worksheet and one Summary"):
        M.choose_sheets(["GRN Report Raw June 26", "Pivot", "Summary"], "x.xlsb")
    with pytest.raises(BuildError, match="more than one 'WS Final' sheet"):
        M.choose_sheets(["WS Final Sheet", "WS Final 2", "Summary"], "x.xlsb")
    assert M.choose_sheets(["GRN Report July 26 ", "Out of formulary July 26 WS", "Pivot", "Summary ",
                            "Need To check flaging later"], "x.xlsb") == \
        ("Out of formulary July 26 WS", "Need To check flaging later", "Summary ")


def test_xlsb_worksheet_rows():
    # pyxlsb: dates and codes as floats, error cells as '0x2a', a price typed as text
    rows = [["OUT OF FORMULARY PURCHASE July 2026"], WS_HEADER,
            W(serial(date(2026, 7, 3)), "HC05", 205411, 11.0, 4.0, 227432, 1067.84, "1,067.84"),
            W(serial(date(2026, 7, 4)), "HC05", 200030, 2.0, 7.0)[:23] + ["0x2a"] * 5 + [None] * 5,
            [None] * 33]
    lines = M.parse_ws_rows(rows, "June WS.xlsb", "WS Final Sheet")
    assert list(lines.columns) == M.WS_LINE_COLUMNS and len(lines) == 2
    assert lines["gr_date"].tolist() == [date(2026, 7, 3), date(2026, 7, 4)]
    assert lines.iloc[0][["material", "suggested", "sug_price", "sug_mrp"]].tolist() == \
        ["205411", "227432", 1067.84, 1067.84]
    assert lines.iloc[1][["suggested", "suggested_desc", "sug_price"]].isna().all()
    with pytest.raises(BuildError, match="no header row"):
        M.parse_ws_rows([["nothing"]] * 3, "x.xlsb", "WS")


def test_xlsb_summary_rows():
    rows = [["ACTUAL OUT OF FORMULARY PURCHASE VALUE June 2026"], SUMMARY_HEADER,
            ["AH01", "Borivalli Mumbai", 287734.56, 508780.53, 221045.97, 0.43, None, 326067.31,
             792036.1, 465968.79, 0.59, 244922.82],
            [None] * 12,
            ["HC17", "Ranchi", 12820.5, 22313.52, 9493.02, 0.43, None, 7982.98, 15565.87, 7582.89,
             0.49, -1910.13],                                     # a negative loss is valid
            ["Grand Total", None, 1.0, 2.0, None, None, None, 3.0, 4.0]]
    s = M.parse_summary_rows(rows, "June WS.xlsb", "Summary")
    assert s.frame["loss"].tolist() == [244922.82, -1910.13] and s.computed == 0
    assert s.titles == ["ACTUAL OUT OF FORMULARY PURCHASE VALUE June 2026"]
    assert len(s.warnings) == 4 and "Grand Total used_pp 1.00" in s.warnings[0]
    with pytest.raises(BuildError, match=r"centres listed twice: \['AH01'\]"):
        M.parse_summary_rows(rows[:3] + rows[2:3], "x.xlsb", "Summary")


# --- the month's Item Master and the Main Sheet (the formulary and spend readers) ----------
def test_item_master_of_the_month(tmp_path):
    dash = tmp_path / "raw" / HASH / DASH
    (dash / "Item Master Sheet").mkdir(parents=True)
    write_master(dash / "Item Master Sheet" / "Item master for formulary complience August -26.xlsx")
    raw = tmp_path / "raw"
    assert list(M.item_master_paths(raw)) == ["2026-08"]
    im = M.load_month_item_master("2026-08", raw)
    f = M.item_master_frame(im).set_index("material")
    assert list(f.reset_index().columns) == M.ITEM_MASTER_COLUMNS
    assert records(f.loc[["101002", "201002"], ["formulary_class", "formulary_flag", "tier"]]) == [
        ["OUT OF FORMULARY", "OUT OF FORMULARY", None], ["FORMULARYP-OPD", "FORMULARY", "P-OPD"]]
    assert f.loc["101001", "minor_group"] == "MINOR 101001"
    with pytest.raises(BuildError, match="no Item Master for 2026-09"):
        M.load_month_item_master("2026-09", raw)


def test_main_sheet_cash_list_prices_and_blocks(tmp_path):
    dash = tmp_path / "raw" / HASH / DASH
    dash.mkdir(parents=True)
    write_scheme_list(dash / "Copy of Onco Formulary -Spend Analysis Main Sheet_11.08.26.xlsx")
    ms = M.load_main_sheet(raw=tmp_path / "raw")
    with_sug = [c for c in CASH if c[6] and c[3] is not None]
    assert sorted(ms.cash.switches) == sorted(str(c[0]) for c in with_sug)
    assert ms.prices["101001"].price == 1000.0 and ms.prices["101003"].gst == 5
    assert "201004" not in ms.prices                       # no GST: not a priced brand
    assert {"material", "block_id", "sharing", "p_rank", "tie"} <= set(ms.blocks.columns)
    assert ms.cash.covers("2026-08")
