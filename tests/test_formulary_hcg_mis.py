"""HCG's PAN formulary-compliance MIS workbook -> reference tables (app/etl/formulary/hcg_mis.py),
and the unit budgets the formulary API takes from them.

The workbook here is built in the test with HCG's real layout and every quirk the real
Sep-26 MIS has: the Apr-26 non-onco OOF repeating the onco value in the FY sheet, unit
actuals on two formulas, an unmarked % block, a label with the wrong year, a block pasted
twice, and a month (Jul-26) that only the FY sheet has.
"""
from datetime import datetime

import openpyxl
import pandas as pd
import pytest

from app.api import formulary
from app.core.config import settings
from app.etl.formulary import hcg_mis
from app.etl.formulary.discover import BuildError
from tests.test_formulary_api import AUG, JUL, overview, write_hospital_map, write_tables

DASH = "Formulary Compliance and Spend Analysis Dashboard data"
FY_SHEET = "FY 26-27 Formulary Compliance %"
OOF = "Approved-Out of formulary consumption"

# unit -> month -> (onco F, onco OOF, non-onco F, non-onco OOF) as the FY sheet prints them
FY_VALUES = {
    "HC05": {"2026-04": (1000.0, 20.0, 100.0, 20.0),     # non-onco OOF repeats onco (true: 5)
             "2026-07": (2000.0, 10.0, 200.0, 30.0)},
    "SH01": {"2026-04": (500.0, 0.0, 300.0, 100.0),
             "2026-07": (600.0, 0.0, 300.0, 60.0)},
}
BUDGET = {"HC05": 1.0, "SH01": 0.92}


def _put(ws, row, col, value):
    ws.cell(row=row, column=col + 1, value=value)


def _fy_sheet(wb):
    ws = wb.create_sheet(FY_SHEET)
    _put(ws, 2, 1, "PAN HCG Formulary Compliance % FY 26-27")
    _put(ws, 3, 1, "Unit Code")
    _put(ws, 3, 2, "Unit Name ")
    col = 3
    for period, extra in (("2026-04", False), ("2026-07", True)):
        _put(ws, 3, col, "Onco ")
        _put(ws, 3, col + 2, "Non Onco")
        heads = ["Formulary consumption", OOF, "Formulary consumption", OOF,
                 "Onco & Non Onco Formulary Consumption ", "Onco & Non Onco Approved-Out of formulary consumption"]
        if extra:
            heads.append("Onco and Non Onco Formulary & Approved Out of formulary Consumption")
        heads += ["Budget FY 26-27", "Actual "]
        for i, h in enumerate(heads):
            _put(ws, 4, col + i, h)
        _put(ws, 3, col + len(heads) - 2, datetime(int(period[:4]), int(period[5:]), 1))
        for r, (plant, months) in enumerate(FY_VALUES.items(), start=5):
            of, oo, nf, no = months[period]
            f, o = of + nf, oo + no
            actual = f / (f + o) if period == "2026-07" else 1 - o / f      # HCG's two formulas
            vals = [of, oo, nf, no, f, o] + ([f + o] if extra else []) + [BUDGET[plant], actual]
            for i, v in enumerate(vals):
                _put(ws, r, col + i, v)
        f_all = sum(m[period][0] + m[period][2] for m in FY_VALUES.values())
        o_all = sum(m[period][1] + m[period][3] for m in FY_VALUES.values())
        _put(ws, 7, col + len(heads) - 1, f_all / (f_all + o_all))
        col += len(heads)
    _put(ws, 4, col, "Remarks -PAN HCG Formulary Compliance Overview for the period April'26 to Sep '26")
    for r, plant in enumerate(FY_VALUES, start=5):
        _put(ws, r, 1, plant)
        _put(ws, r, 2, {"HC05": "KR BANGALORE", "SH01": "SUCHIRAYU-HUBLI"}[plant])
        _put(ws, r, col, {"HC05": "Compliance maintained", "SH01": "Multispecialty unit; improvement observed"}[plant])
    _put(ws, 7, 1, "Overall Onco and Non Onco Formulary Compliance % FY 26-27")


def _history_sheet(wb, name, code_col, blocks, units):
    """blocks: [(label, [headers], {plant: [values]})] in column order."""
    ws = wb.create_sheet(name)
    _put(ws, 1, 0, f"{name} compliances")
    for i, h in enumerate(["Code", "Region", "Plant name"]):
        _put(ws, 2, code_col + i, h)
    for r, (plant, region) in enumerate(units, start=3):
        _put(ws, r, code_col, plant)
        _put(ws, r, code_col + 1, region)
        _put(ws, r, code_col + 2, plant.lower())
    _put(ws, 3 + len(units), code_col + 2, "Average")
    col = code_col + 4
    for label, heads, vals in blocks:
        _put(ws, 1, col, label)
        for i, h in enumerate(heads):
            _put(ws, 2, col + i, h)
        for r, (plant, _) in enumerate(units, start=3):
            for i, v in enumerate(vals.get(plant, [])):
                _put(ws, r, col + i, v)
        col += len(heads) + 1


VALUE = ["Formulary consumption", OOF, "Non formulary consumption", "Restricted brand consumption for scheme patients"]


def write_mis(folder, name="Onco and Non Onco Formulary Compliance Sep 26 - MIS.xlsx", extra_onco_block=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _fy_sheet(wb)
    units = [("HC05", "Karnataka "), ("SH01", "Karnataka "), ("HC03", "Karnataka ")]
    april = ("April 23 in Value", VALUE, {"HC05": [1000, 20, 7, 50], "SH01": [500, 0, 3, 0], "HC03": [10, 1, 0, 0]})
    onco = [
        (datetime(2026, 3, 1), ["Formulary consumption", OOF], {"HC05": [0.95, 0.05], "SH01": [0.9, 0.1]}),   # unmarked %
        ("March 2026 in Value", VALUE, {"HC05": [900, 30, 5, 40], "SH01": [450, 10, 2, 0]}),
        ("April 2026 in %", ["Formulary consumption in %", f"{OOF} in %"], {"HC05": [0.98, 0.02]}),
        april,                     # 2-digit year that is wrong: Apr-26 by its place
        april,                     # pasted twice
        ("Sep 2026 in Value", VALUE, {"HC05": [1100, 15, 6, 45], "SH01": [520, 5, 1, 0]}),
    ] + ([extra_onco_block] if extra_onco_block else [])
    non = [
        ("April 2026 in Value", VALUE, {"HC05": [100, 5, 9, 1], "SH01": [300, 100, 4, 0], "HC03": [5, 0, 0, 0]}),
        ("Sep 2026 in Value", VALUE, {"HC05": [110, 4, 8, 1], "SH01": [310, 50, 3, 0]}),
    ]
    _history_sheet(wb, "Onco Formulary Consumption", 1, onco, units)
    _history_sheet(wb, "Non Onco Formulary Consumption", 0, non, units)
    d = folder / "raw" / "x" / DASH / "HCG MIS"
    d.mkdir(parents=True, exist_ok=True)
    wb.save(d / name)
    return folder / "raw"


@pytest.fixture()
def built(tmp_path):
    raw = write_mis(tmp_path)
    report = hcg_mis.build(raw=raw, data=tmp_path / "data", log=lambda *_: None)
    tables = {n: pd.read_parquet(tmp_path / "data" / f"{n}.parquet") for n in hcg_mis.C.MIS_TABLES}
    return report, tables


def test_history_corrects_the_fy_sheet_and_formulas_are_detected(built):
    report, t = built
    m = t["hcg_mis_monthly"].set_index(["plant", "period", "segment"])
    # Apr-26: the FY sheet's repeated non-onco OOF (20) is replaced by history's 5
    assert m.loc[("HC05", "2026-04", "non-onco"), "out_of_formulary"] == 5
    assert m.loc[("HC05", "2026-04", "non-onco"), "source"] == "history"
    assert [c["reason"] for c in report["corrections"]] == ["non-onco OOF repeats the onco OOF in the FY sheet"]
    # Jul-26 has no history block: the FY sheet's rupees, nothing else
    jul = m.loc[("HC05", "2026-07", "onco")]
    assert jul["source"] == "fy_sheet" and jul["formulary"] == 2000 and pd.isna(jul["non_formulary"])
    c = t["hcg_mis_compliance"].set_index(["plant", "period"])
    assert c.loc[("HC05", "2026-04"), "budget_pct"] == 100 and c.loc[("SH01", "2026-07"), "budget_pct"] == 92
    assert c.loc[("HC05", "2026-04"), "hcg_formula"] == "oof_over_formulary"
    assert c.loc[("HC05", "2026-07"), "hcg_formula"] == "pooled"
    assert c.loc[("HC05", "2026-04"), "hcg_actual_pct"] == pytest.approx(100 * (1 - 40 / 1100))
    assert c.loc[("HC05", "2026-04"), "pooled_pct"] == pytest.approx(100 * 1100 / 1125)


def test_pan_row_units_and_regions(built):
    report, t = built
    pan = t["hcg_mis_pan"].set_index("period")
    assert pan.loc["2026-04", "printed_pct"] == pytest.approx(100 * 1900 / 2040)
    assert pan.loc["2026-04", "sheet_pooled_pct"] == pytest.approx(100 * 1900 / 2040)
    assert pan.loc["2026-04", "corrected_pooled_pct"] == pytest.approx(100 * 1900 / 2025)
    u = t["hcg_mis_units"].set_index("plant")
    assert list(t["hcg_mis_units"]["plant"]) == ["HC05", "SH01", "HC03"]
    assert u.loc["SH01", "remark"] == "Multispecialty unit; improvement observed" and u.loc["SH01", "region"] == "Karnataka"
    assert not u.loc["HC03", "in_pan"] and pd.isna(u.loc["HC03", "pan_order"])
    assert report["history"]["onco"] == {"first": "2026-03", "last": "2026-09",
                                         "missing_inside": ["2026-05", "2026-06", "2026-07", "2026-08"]}


def test_month_labels_resolved_and_repeats_read_once(built):
    report, t = built
    notes = "\n".join(report["notes"])
    assert "unmarked block '2026-03-01 00:00:00' read as percent" in notes
    assert "'April 23 in Value' read as 2026-04 (its label says 2023)" in notes
    assert "repeats the one at column" in notes
    m = t["hcg_mis_monthly"]
    assert len(m[(m["plant"] == "HC05") & (m["period"] == "2026-04") & (m["segment"] == "onco")]) == 1
    assert m[(m["plant"] == "HC05") & (m["period"] == "2026-03")]["formulary"].tolist() == [900]


def test_unknown_column_aborts_and_nothing_is_published(tmp_path):
    raw = write_mis(tmp_path, extra_onco_block=("Oct 2026 in Value", ["Formulary consumption", "Something new"],
                                                {"HC05": [1, 2]}))
    with pytest.raises(BuildError, match="unknown column 'Something new'"):
        hcg_mis.build(raw=raw, data=tmp_path / "data", log=lambda *_: None)
    assert not (tmp_path / "data" / "hcg_mis_units.parquet").exists()


def test_latest_workbook_is_read_and_two_for_one_month_abort(tmp_path):
    write_mis(tmp_path, name="Onco and Non Onco Formulary Compliance Aug 26 - MIS.xlsx")
    raw = write_mis(tmp_path)
    latest, every = hcg_mis.find_mis(raw)
    assert latest.period == "2026-09" and [m.period for m in every] == ["2026-08", "2026-09"]
    write_mis(tmp_path, name="Onco and Non Onco Formulary Compliance Sep 2026 - MIS.xlsx")
    with pytest.raises(BuildError, match="two MIS workbooks for 2026-09"):
        hcg_mis.find_mis(raw)


# ---------- the API's unit budgets ----------

@pytest.fixture()
def data(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "FORMULARY_DATA_DIR", str(tmp_path / "formulary"))
    monkeypatch.setattr(settings, "REFERENCE_DIR", str(tmp_path / "reference"))
    write_tables(tmp_path / "formulary")
    write_hospital_map(tmp_path / "reference")
    formulary.clear_caches()
    yield tmp_path
    formulary.clear_caches()


def _write_budgets(folder):
    pd.DataFrame([
        {"plant": "HC05", "period": "2026-04", "budget_pct": 100.0, "hcg_actual_pct": 97.0, "hcg_formula": "pooled", "pooled_pct": 97.0},
        {"plant": "HS01", "period": "2026-04", "budget_pct": 85.0, "hcg_actual_pct": 80.0, "hcg_formula": "pooled", "pooled_pct": 80.0},
        {"plant": "HS01", "period": JUL, "budget_pct": 85.0, "hcg_actual_pct": 79.0, "hcg_formula": "pooled", "pooled_pct": 79.0},
    ]).to_parquet(folder / "hcg_mis_compliance.parquet", index=False)


def test_budget_defaults_until_the_mis_is_read(client, data):
    h = overview(client, Plant="HS01", Month=AUG)["headline"]
    assert h["budget_pct"] == 100 and h["budget_source"] == "default"
    assert overview(client, Plant="ALL", Month=AUG)["headline"]["budget_pct"] is None


def test_unit_budgets_from_the_mis(client, data):
    _write_budgets(data / "formulary")
    formulary.clear_caches()
    o = overview(client, Plant="HS01", Month=AUG)
    # Aug is not in the MIS rows here: the unit's FY 26-27 budget applies
    assert o["headline"]["budget_pct"] == 85 and o["headline"]["budget_source"] == "hcg_mis"
    assert {t["period"]: t["budget_pct"] for t in o["trend"]} == {JUL: 85, AUG: 85}
    # a unit HCG sets no budget for, and all hospitals together, have none
    assert overview(client, Plant="HC06", Month=AUG)["headline"]["budget_pct"] is None
    a = overview(client, Plant="ALL", Month=AUG)
    assert a["headline"]["budget_pct"] is None
    assert {b["plant"]: b["budget_pct"] for b in a["by_hospital"]}["HS01"] == 85
